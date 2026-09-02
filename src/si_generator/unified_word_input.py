from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
import os
from pathlib import Path
import re
import shutil
import tempfile
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.shared import Pt
from lxml import etree


SECTION_MARKERS = {
    "compound_table": "[AUTO SI: COMPOUND TABLE]",
    "reaction_schema": "[AUTO SI: REACTION SCHEMA]",
    "scope": "[AUTO SI: SCOPE]",
    "si_template": "[AUTO SI: SI TEMPLATE]",
}
END_MARKER = "[AUTO SI: END]"


class UnifiedInputError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class UnifiedInputBundle:
    source: Path
    compound_table: Path
    reaction_schema: Path | None = None
    scope: Path | None = None
    si_template: Path | None = None

    @property
    def has_complete_loadings(self) -> bool:
        return bool(self.reaction_schema and self.scope)


def materialize_unified_input(source: str | Path, output_dir: str | Path) -> UnifiedInputBundle:
    source_path = Path(source).expanduser().resolve()
    if not source_path.exists() or not source_path.is_file():
        raise UnifiedInputError(f"All-in-one input DOCX does not exist: {source_path}")
    if source_path.suffix.lower() != ".docx" or not zipfile.is_zipfile(source_path):
        raise UnifiedInputError(f"All-in-one input must be a readable .docx file: {source_path}")

    sections = _read_sections(source_path)
    compound_elements = sections.get("compound_table", [])
    if not compound_elements or not any(element.tag == qn("w:tbl") for element in compound_elements):
        raise UnifiedInputError(
            f"All-in-one input must contain {SECTION_MARKERS['compound_table']} followed by the compound table."
        )

    target_root = Path(output_dir).expanduser().resolve()
    if target_root.exists():
        shutil.rmtree(target_root)
    target_root.mkdir(parents=True, exist_ok=True)

    paths: dict[str, Path | None] = {}
    filenames = {
        "compound_table": "Compound_table.docx",
        "reaction_schema": "Reaction_schema.docx",
        "scope": "Scope.docx",
        "si_template": "SI_template.docx",
    }
    for key, filename in filenames.items():
        elements = sections.get(key, [])
        if not _has_meaningful_content(elements):
            paths[key] = None
            continue
        target = target_root / filename
        _write_section_document(source_path, target, elements)
        paths[key] = target

    return UnifiedInputBundle(
        source=source_path,
        compound_table=paths["compound_table"],
        reaction_schema=paths["reaction_schema"],
        scope=paths["scope"],
        si_template=paths["si_template"],
    )


def default_unified_staging_dir(run_id: str) -> Path:
    base = Path(os.environ.get("LOCALAPPDATA") or tempfile.gettempdir()) / "AutoSupportGenerator" / "staging"
    safe_run_id = "".join(character for character in str(run_id) if character.isalnum() or character in "-_") or "run"
    return base / f"unified_{safe_run_id}"


def build_unified_input_docx(
    compound_table: str | Path,
    output_path: str | Path,
    *,
    reaction_schema: str | Path | None = None,
    scope: str | Path | None = None,
    si_template: str | Path | None = None,
) -> Path:
    compound_path = Path(compound_table).resolve()
    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    base_document = Path(si_template).resolve() if si_template else compound_path
    shutil.copy2(base_document, output)

    document = Document(output)
    body = document._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    title = document.add_paragraph("Auto Support Generator - all-in-one input")
    title_run = title.runs[0]
    title_run.bold = True
    title_run.font.size = Pt(18)
    note = document.add_paragraph(
        "Edit the tables and template below. Keep the section labels unchanged. "
        "Reaction schema, Scope and SI template are optional."
    )
    marker = document.add_paragraph(SECTION_MARKERS["compound_table"])
    _format_section_marker(marker)
    section_sources = (
        ("compound_table", compound_path),
        ("reaction_schema", reaction_schema),
        ("scope", scope),
        ("si_template", si_template),
    )
    for index, (key, source) in enumerate(section_sources):
        if not source:
            continue
        if index:
            heading = document.add_paragraph(SECTION_MARKERS[key])
            _format_section_marker(heading)
            heading.paragraph_format.page_break_before = True
        source_document = Document(Path(source).resolve())
        for element in source_document._element.body:
            if element.tag == qn("w:sectPr"):
                continue
            copied = deepcopy(element)
            _remap_element_relationships(copied, source_document.part, document.part)
            body.insert(len(body) - 1, copied)

    end = document.add_paragraph(END_MARKER)
    _format_section_marker(end)
    document.save(output)
    return output


def _format_section_marker(paragraph) -> None:
    if paragraph.runs:
        paragraph.runs[0].bold = True
        paragraph.runs[0].font.size = Pt(13)


def _remap_element_relationships(element, source_part, target_part) -> None:
    cloned_parts: dict[str, Part] = {}
    relationship_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for node in element.iter():
        for attribute_name, relationship_id in list(node.attrib.items()):
            if not attribute_name.startswith("{" + relationship_namespace + "}"):
                continue
            relationship = source_part.rels.get(relationship_id)
            if relationship is None:
                continue
            if relationship.is_external:
                new_id = target_part.relate_to(relationship.target_ref, relationship.reltype, is_external=True)
            else:
                source_key = str(relationship.target_part.partname)
                cloned_part = cloned_parts.get(source_key)
                if cloned_part is None:
                    cloned_part = _clone_part(relationship.target_part, target_part.package)
                    cloned_parts[source_key] = cloned_part
                new_id = target_part.relate_to(cloned_part, relationship.reltype)
            node.set(attribute_name, new_id)


def _clone_part(source_part, target_package) -> Part:
    source_name = str(source_part.partname)
    match = re.match(r"^(.*?)(\d+)(\.[^./]+)$", source_name)
    if match:
        template = f"{match.group(1)}%d{match.group(3)}"
    else:
        stem, suffix = source_name.rsplit(".", 1) if "." in source_name else (source_name, "bin")
        template = f"{stem}_%d.{suffix}"
    partname = target_package.next_partname(template)
    cloned = Part(partname, source_part.content_type, source_part.blob, target_package)
    for relationship in source_part.rels.values():
        if relationship.is_external:
            cloned.relate_to(relationship.target_ref, relationship.reltype, is_external=True)
        else:
            cloned.relate_to(_clone_part(relationship.target_part, target_package), relationship.reltype)
    return cloned


def _read_sections(source: Path) -> dict[str, list[etree._Element]]:
    with zipfile.ZipFile(source, "r") as archive:
        try:
            root = etree.fromstring(archive.read("word/document.xml"))
        except (KeyError, etree.XMLSyntaxError) as exc:
            raise UnifiedInputError(f"Cannot read word/document.xml from {source}: {exc}") from exc

    body = root.find(qn("w:body"))
    if body is None:
        raise UnifiedInputError(f"All-in-one input has no Word document body: {source}")

    marker_to_key = {_normalize_marker(value): key for key, value in SECTION_MARKERS.items()}
    sections: dict[str, list[etree._Element]] = {}
    current_key: str | None = None
    for element in body:
        if element.tag == qn("w:sectPr"):
            continue
        if element.tag == qn("w:p"):
            marker_text = _normalize_marker(_element_text(element))
            if marker_text == _normalize_marker(END_MARKER):
                current_key = None
                continue
            if marker_text in marker_to_key:
                current_key = marker_to_key[marker_text]
                if current_key in sections:
                    raise UnifiedInputError(f"Duplicate section marker: {SECTION_MARKERS[current_key]}")
                sections[current_key] = []
                continue
        if current_key:
            sections[current_key].append(deepcopy(element))
    return sections


def _write_section_document(source: Path, target: Path, elements: list[etree._Element]) -> None:
    with zipfile.ZipFile(source, "r") as input_archive:
        root = etree.fromstring(input_archive.read("word/document.xml"))
        body = root.find(qn("w:body"))
        if body is None:
            raise UnifiedInputError(f"All-in-one input has no Word document body: {source}")
        section_properties = deepcopy(body.find(qn("w:sectPr")))
        for child in list(body):
            body.remove(child)
        for element in elements:
            body.append(deepcopy(element))
        if section_properties is not None:
            body.append(section_properties)
        document_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output_archive:
            for item in input_archive.infolist():
                payload = document_xml if item.filename == "word/document.xml" else input_archive.read(item.filename)
                output_archive.writestr(item, payload)


def _has_meaningful_content(elements: list[etree._Element]) -> bool:
    return any(element.tag == qn("w:tbl") or _element_text(element).strip() for element in elements)


def _element_text(element: etree._Element) -> str:
    return "".join(element.itertext())


def _normalize_marker(value: str) -> str:
    return " ".join(str(value).replace("_", " ").split()).upper()
