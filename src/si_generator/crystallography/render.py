from __future__ import annotations

from copy import deepcopy
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.image.image import Image

from .fields import chemical_formula_runs
from .numbering import SEQUENCE, renumber_tables


def load_data(compound):
    if not compound.crystallography_data_path:
        return None, {}
    path = Path(compound.crystallography_data_path)
    return path.parent, json.loads(path.read_text(encoding="utf-8"))


def append_figures(document, compound):
    base, data = load_data(compound)
    for record in data.get("records", []):
        p = document.add_paragraph()
        p.paragraph_format.keep_with_next = True
        section = document.sections[-1]
        width = min(Inches(4.8), section.page_width - section.left_margin - section.right_margin)
        image_path = base / record["image"]
        img = Image.from_file(str(image_path))
        width = min(width, int(Inches(3.2) * img.px_width / img.px_height))
        p.add_run().add_picture(str(image_path), width=width)
        caption = document.add_paragraph("Crystal structure of ")
        caption.add_run(compound.number)
        caption.add_run(". " + record['caption'].replace('{Product.number}', compound.number))
        caption.paragraph_format.keep_together = True
        if record["ccdc"]:
            caption.add_run(f" Deposition: {record['ccdc']}.")


def _numbered_paragraph(document, text, number):
    text = text.replace("{Product.number}", number)
    p = document.add_paragraph()
    for piece in re.split(r"(?<![A-Za-z0-9])(" + re.escape(number) + r")(?![A-Za-z0-9])", text):
        p.add_run(piece)
    return p


def _caption(document, text, number):
    p = document.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(5)
    p.add_run("Table S").bold = True
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"SEQ {SEQUENCE} \\* ARABIC")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    properties.append(OxmlElement("w:b"))
    run.append(properties)
    value = OxmlElement("w:t")
    value.text = "1"
    run.append(value)
    field.append(run)
    p._p.append(field)
    p.add_run(".").bold = True
    p.add_run(f" {text} for compound ")
    p.add_run(number).bold = True
    p.add_run(".")
    renumber_tables(document.element)


def _table(document, rows, *, headers=None, formula_row=None):
    columns = len(headers) if headers else 2
    table = document.add_table(rows=0, cols=columns)
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    section = document.sections[-1]
    width = min(Inches(5.35), section.page_width - section.left_margin - section.right_margin)
    widths = [int(width * ratio) for ratio in ((.53, .47) if columns == 2 else (.15, .35, .30, .20))]
    for column, column_width in zip(table.columns, widths):
        column.width = column_width
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "808080")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    values_to_render = ([headers] if headers else []) + list(rows)
    for row_index, values in enumerate(values_to_render):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if headers and row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column_index, (cell, value, column_width) in enumerate(zip(row.cells, values, widths)):
            cell.width = column_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == formula_row and column_index == 1:
                for text, subscript in chemical_formula_runs(str(value)):
                    run = cell.paragraphs[0].add_run(text)
                    run.font.subscript = subscript
            else:
                cell.text = str(value)
            for run in cell.paragraphs[0].runs:
                run.bold = bool(headers and row_index == 0)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.left_indent = Pt(0)
                p.paragraph_format.right_indent = Pt(0)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.keep_with_next = False
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
    document.add_paragraph()


def append_details(document, compound):
    base, data = load_data(compound)
    template = Document(base / data["template"]) if data.get("template") else None
    for record in data.get("records", []):
        values = {"Product.number": compound.number, "Product.name": compound.display_name,
                  "Crystal.block": record["block"], "Crystal.ccdc": record["ccdc"]}
        paragraphs = template.paragraphs if template else None
        text_lines = [p.text for p in paragraphs] if paragraphs is not None else [
            "Crystallographic data for {Product.number}", "{Crystal.description}", "{Crystal.table}", "{Crystal.geometry}"]
        if template and "{Crystal.table}" not in text_lines:
            raise ValueError("Crystallography template must contain {Crystal.table} in a separate paragraph")
        for i, text in enumerate(text_lines):
            if text == "{Crystal.description}":
                for description in record["description"]:
                    _numbered_paragraph(document, description, compound.number)
                if record["ccdc"]:
                    document.add_paragraph(f"Crystallographic deposition: {record['ccdc']}.")
            elif text == "{Crystal.table}":
                _caption(document, "Crystal data and structure refinement", compound.number)
                rows = [("Identification code", f"CCDC {record['ccdc']}" if record["ccdc"] else record["block"])]
                formula_row = None
                for row in record["rows"]:
                    if row["value"] is not None:
                        if row["key"] == "formula":
                            formula_row = len(rows)
                        rows.append((row["label"], row["value"]))
                _table(document, rows, formula_row=formula_row)
            elif text == "{Crystal.geometry}":
                if record["geometry"]:
                    _caption(document, "Bond lengths and angles", compound.number)
                    _table(document, [(g["kind"], g["atoms"], g["value"], g["symmetry"]) for g in record["geometry"]],
                           headers=["Geometry", "Atoms", "Value (angstrom / degrees)", "Symmetry"])
            else:
                for key, value in values.items():
                    if key == "Product.number":
                        continue
                    text = text.replace("{" + key + "}", value)
                p = document.add_paragraph()
                for part_index, part in enumerate(text.split("{Product.number}")):
                    if part_index:
                        p.add_run(compound.number)
                    p.add_run(part)
                if template and paragraphs[i]._p.pPr is not None:
                    p._p.insert(0, deepcopy(paragraphs[i]._p.pPr))
                if i == 0:
                    p.paragraph_format.keep_with_next = True
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0, 0, 0)


def write_report(compound, target):
    document = Document()
    for border in document.styles['Title'].element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    title = document.add_heading(f"X-ray crystallography of {compound.number}", 0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    append_figures(document, compound)
    document.add_page_break()
    append_details(document, compound)
    document.save(target)
