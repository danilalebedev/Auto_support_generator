from __future__ import annotations

import csv
import json
import re
import shutil
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor

from .reagent_catalog import lookup_reagent, normalize_reagent_name, catalog_summary
from .unified_word_input import build_unified_input_docx


NUMBER = r"(?:\d+\s*/\s*[1-9]\d*|\d+(?:[.,]\d+)?)"
UNIT = (
    r"(?:µL|μL|uL|mL|L|µmol|μmol|umol|mmol|mol\s*%|mol|µg|μg|ug|mg|kg|g|"
    r"mM|M|equivalents?|equiv\.?|eq\.?|wt\s*%|%)"
)
# mol/L must precede mol so that a stock concentration is not read as moles.
UNIT = r"(?:mol\s*/\s*L|" + UNIT + ")"
QUANTITY_RE = re.compile(rf"(?<![\d.])(?<!\d/)(?P<value>{NUMBER})\s*(?P<unit>{UNIT})(?![A-Za-z])", re.IGNORECASE)
PAREN_RE = re.compile(r"\((?P<body>[^()]{1,200})\)")

SCHEMA_HEADERS = ("Reagents", "equiv.", "MW, g/mol", "Density, g/ml", "Concentration, M")
LOADINGS_HEADERS = (
    "Alias",
    "Compound",
    "Role / scope",
    "Source loading",
    "Mass, mg",
    "Amount, mmol",
    "Volume, mL",
    "equiv.",
    "Concentration, M",
    "MW, g/mol",
    "Density, g/mL",
    "Source / review",
)
LOADINGS_COLUMN_WIDTHS_INCHES = (0.68, 1.32, 0.78, 1.12, 0.54, 0.58, 0.58, 0.46, 0.67, 0.63, 0.68, 1.72)

SOLVENT_NAMES = {
    "acetonitrile",
    "acn",
    "benzene",
    "ch2cl2",
    "chcl3",
    "chloroform",
    "dichloromethane",
    "dcm",
    "dioxane",
    "dmf",
    "dmso",
    "ethanol",
    "etoh",
    "ethyl acetate",
    "etoac",
    "ether",
    "hexane",
    "mech2cl",
    "mecn",
    "methanol",
    "meoh",
    "mtbe",
    "n-hexane",
    "petroleum ether",
    "t-butanol",
    "tert-butanol",
    "tetrahydrofuran",
    "thf",
    "toluene",
    "water",
    "1,2-dichloroethane", "1,4-dioxane", "2-methyltetrahydrofuran", "dce", "et2o",
    "hfip", "tfe", "2,2,2-tfe", "c6hf5", "c6f5cf3", "o-xylene", "nmp", "dma",
    "ch3cn", "h2o", "dme", "1,2-dimethoxyethane", "solvent",
    "isopropanol", "iproh", "i-proh", "2-propanol", "propan-2-ol", "phch3", "phme",
    "diethyl ether", "acetone", "hexanes", "heptane", "decane", "cyclohexane", "diglyme",
    "dichloroethane", "N,N-dimethylformamide", "dimethylformamide", "dimethyl sulfoxide",
    "1,1,1,3,3,3-hexafluoro-2-propanol", "hexafluoroisopropanol", "2,2,2-trifluoroethanol",
}
WORKUP_CONTEXT_RE = re.compile(
    r"\b(?:quenched|quench|extracted|extract|washed|wash|diluted|dilute|partitioned|"
    r"filtered|filter|filtration|rinsed|rinse|purified|chromatography|recrystalli[sz](?:ed|ation)|triturated)\b",
    re.IGNORECASE,
)


@dataclass
class ParsedQuantity:
    kind: str
    value: float
    canonical_unit: str
    source_unit: str
    raw: str
    span: tuple[int, int]
    derived: bool = False
    formula: str | None = None


@dataclass
class ParsedChemical:
    name: str
    role: str = "reagent"
    alias: str = ""
    variable: bool = False
    name_spans: list[tuple[int, int]] = field(default_factory=list)
    quantities: list[ParsedQuantity] = field(default_factory=list)
    molecular_weight_g_mol: float | None = None
    density_g_ml: float | None = None
    warnings: list[str] = field(default_factory=list)
    property_sources: dict[str, Any] = field(default_factory=dict)
    catalog_status: str = "not_found"
    formulation: bool = False

    def value(self, kind: str) -> float | None:
        return next((quantity.value for quantity in self.quantities if quantity.kind == kind), None)

    def source_quantity(self, kind: str) -> ParsedQuantity | None:
        return next(
            (quantity for quantity in self.quantities if quantity.kind == kind and not quantity.derived),
            None,
        )

    def add_quantity(self, quantity: ParsedQuantity) -> None:
        if any(existing.kind == quantity.kind and existing.span == quantity.span for existing in self.quantities):
            return
        self.quantities.append(quantity)


@dataclass(frozen=True)
class GeneratedProcedureInputs:
    compound_table: Path
    reaction_schema: Path
    scope: Path
    si_template: Path
    all_in_one: Path
    report: Path
    loadings_table: Path
    source_method: Path | None = None

    @property
    def scope_draft(self) -> Path:
        """Compatibility alias for callers created before Scope.docx became final output."""
        return self.scope


def read_procedure_text(path: str | Path) -> str:
    """Read a plain-text, Markdown, or Word procedure into one normalized string."""
    source = Path(path)
    if source.suffix.casefold() == ".docx":
        document = Document(str(source))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = str(getattr(paragraph.style, "name", "") or "").casefold()
            common_heading = bool(
                re.fullmatch(
                    r"(?:general|typical|representative)(?:\s+experimental)?\s+procedure(?:\s+[A-Za-z0-9.-]+)?",
                    text,
                    re.IGNORECASE,
                )
            )
            if style_name.startswith(("title", "heading")) or common_heading:
                continue
            blocks.append(text)
        for table in document.tables:
            for row in table.rows:
                line = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if line:
                    blocks.append(line)
        return "\n".join(blocks)
    return source.read_text(encoding="utf-8-sig")


def parse_procedure(
    text: str,
    *,
    variable_names: Iterable[str] = (),
    inventory_path: str | Path | None = None,
) -> dict[str, Any]:
    """Parse one experimental procedure into an auditable intermediate representation.

    The parser intentionally handles explicit quantities and relationships only. It does
    not invent structures or silently accept ambiguous reagent names.
    """
    normalized_text = _normalize_text(text)
    variable_names = list(variable_names)
    chemicals_by_name: dict[str, ParsedChemical] = {}

    def chemical_for(name: str, span: tuple[int, int], *, context: str = "") -> ParsedChemical | None:
        cleaned = _clean_name(name)
        # A user-supplied exact name disambiguates generic prose prefixes while
        # retaining those prefixes outside the substituted name span.
        for requested in variable_names:
            marked_suffix = (not re.fullmatch(r"\d+[a-z]?",requested,re.I)
                             and _plausible_chemical_name(cleaned) and not QUANTITY_RE.search(cleaned)
                             and re.search(r"(?<![A-Za-z0-9])"+re.escape(_normalize_text(requested))+r"$", cleaned, re.I))
            if _same_chemical_name(cleaned, requested) or marked_suffix:
                mention = re.search(re.escape(_normalize_text(requested)), normalized_text[span[0]:span[1]], re.I)
                if mention:
                    span = (span[0]+mention.start(), span[0]+mention.end())
                    cleaned = normalized_text[span[0]:span[1]]
                    break
        marked_short_label = bool(re.fullmatch(r"\d+[a-z]", cleaned, re.I)) and any(_same_chemical_name(cleaned,n) for n in variable_names)
        if not _plausible_chemical_name(cleaned) and not marked_short_label:
            return None
        role = _infer_role(cleaned, context)
        marked_after_workup = False
        if _is_workup_at(normalized_text, span[0]) and role != "solvent" and any(_same_chemical_name(cleaned,n) for n in variable_names):
            marked_after_workup = True
        elif _is_workup_at(normalized_text, span[0]) or _is_terminal_workup_addition(normalized_text,span):
            role = "workup"
        # Reaction and workup occurrences of the same solvent are separate.
        identity = _normalize_name(cleaned)
        if role == "solvent":
            identity = normalize_reagent_name(re.sub(r"^anhyd\.?\s+", "anhydrous ", cleaned, flags=re.I))
        key = identity + ("::workup" if role == "workup" else "")
        chemical = chemicals_by_name.get(key)
        if chemical is None:
            chemical = ParsedChemical(name=cleaned, role=role)
            chemicals_by_name[key] = chemical
        if marked_after_workup:
            chemical.warnings.append("Marked substrate appears after workup; stage review required.")
        if span not in chemical.name_spans:
            chemical.name_spans.append(span)
        return chemical

    # Most SI procedures use "name (mass, mmol, equiv.)". Extract those blocks first.
    loading_blocks = list(_loading_parentheses(normalized_text))
    consumed_quantities: list[tuple[int, int]] = []
    for block_start, body_start, block_end in loading_blocks:
        # In "1.0 equiv (0.5 mmol) of substrate" the parenthetical belongs to
        # the following chemical and is handled once by the prefix pattern.
        before = normalized_text[max(0, block_start - 40) : block_start]
        after = normalized_text[block_end : block_end + 40]
        marked_label_before = any(re.search(r"(?<![A-Za-z0-9])"+re.escape(_normalize_text(n))+r"\s*$", before, re.I) for n in variable_names)
        if not marked_label_before and re.search(
            rf"{NUMBER}\s*{UNIT}\s*$",
            before,
            re.IGNORECASE,
        ) and re.match(r"\s+(?!(?:was|were|and|in|at|to|for|then|followed|subsequently)\b)[A-Za-z0-9(+]", after, re.IGNORECASE):
            continue
        body = normalized_text[body_start:block_end-1]
        # Isolated-product results and equipment dimensions are not additions.
        if re.search(r"\byield\b", body, re.I):
            continue
        matches = list(QUANTITY_RE.finditer(body))
        matches = [m for m in matches if not body[m.end():].lstrip().startswith("/")]
        chemistry_matches = [match for match in matches if _quantity_kind(match.group("unit")) in _CHEMISTRY_KINDS]
        if not chemistry_matches:
            continue
        name, name_span, context = _name_left_of(normalized_text, block_start)
        aside = re.search(r",\s*(?:prepared|synthesized|obtained)\s+as\b[^()]{0,100}$", normalized_text[max(0,block_start-130):block_start], re.I)
        if aside:
            name, name_span, context = _name_left_of(normalized_text, max(0,block_start-130)+aside.start())
        # Also accept "substrate 120 mg (0.5 mmol, 1 equiv)". The mass is
        # outside the loading parentheses but belongs to the preceding name.
        suffix_quantity = re.search(rf"\s+(?P<q>{NUMBER}\s*(?:mg|g|kg|mL|µL|uL))\s*$", name, re.I)
        trailing_quantity = None
        if suffix_quantity:
            trailing_quantity = QUANTITY_RE.search(suffix_quantity["q"])
            trailing_offset = name_span[0] + suffix_quantity.start("q")
            name = name[:suffix_quantity.start()].rstrip()
            name_span = (name_span[0], name_span[0]+len(name))
        leading_quantity = re.match(rf"\s*(?P<q>{NUMBER}\s*(?:mg|g|kg|mL|µL|uL))\s+(?P<name>.+)$",name,re.I)
        leading_parsed = None
        if leading_quantity:
            leading_match=QUANTITY_RE.search(leading_quantity["q"])
            leading_offset=name_span[0]+leading_quantity.start("q")
            leading_parsed=_parse_quantity(leading_match,leading_offset)
            name_start=name_span[0]+leading_quantity.start("name")
            name=leading_quantity["name"].strip()
            name_span=(name_start,name_start+len(name))
        mixture = _mixture_left_of(normalized_text, block_start)
        if mixture is not None:
            name, name_span = mixture
        stock = _stock_left_of(normalized_text, block_start, chemistry_matches)
        if stock is not None:
            name, name_span, carrier = stock
        if re.search(r"\bwire\b",name,re.I) and not any(_quantity_kind(m["unit"]) in {"mass_mg","amount_mmol"} for m in chemistry_matches):
            continue
        chemical = chemical_for(name, name_span, context=context)
        if chemical is None:
            continue
        # `DCM/toluene (14.7 mL/14.7 mL)` and `Pd/SPhos (10 mg/12 mg)`
        # describe component allocations which cannot be represented by one
        # flat scalar alias. Preserve the source literal for review.
        paired_spans=[]
        for pair in re.finditer(rf"{NUMBER}\s*(?P<u>mL|mg|g)\s*/\s*{NUMBER}\s*(?P=u)",body,re.I):
            paired_spans.append(pair.span())
        chemistry_matches=[m for m in chemistry_matches if not any(a<=m.start() and m.end()<=b for a,b in paired_spans)]
        if paired_spans:
            chemical.warnings.append("Component-specific slash loadings remain literal; a flat alias would misassign them.")
        if stock is not None:
            chemical.formulation = True
            chemical.property_sources["stock_carrier"] = {"source":"explicit_text","name":carrier}
        inner_carrier = re.search(r"\b(?:solution\s+)?in\s+(?P<name>[A-Za-z0-9/-]+)(?=\s*[,;)]|\s*$)",body,re.I)
        if inner_carrier and any(_quantity_kind(m["unit"])=="concentration_m" for m in chemistry_matches):
            carrier_name = inner_carrier["name"]
            if normalize_reagent_name(carrier_name) in {normalize_reagent_name(n) for n in SOLVENT_NAMES} or _solvent_mixture_parts(carrier_name):
                chemical.formulation = True
                chemical.property_sources["stock_carrier"] = {"source":"explicit_text","name":carrier_name}
        if trailing_quantity:
            parsed_trailing = _parse_quantity(trailing_quantity, trailing_offset)
            chemical.add_quantity(parsed_trailing)
            consumed_quantities.append(parsed_trailing.span)
        if leading_parsed:
            chemical.add_quantity(leading_parsed)
            consumed_quantities.append(leading_parsed.span)
        offset = body_start
        for match in chemistry_matches:
            chemical.add_quantity(_parse_quantity(match, offset))
            consumed_quantities.append((offset+match.start(), offset+match.end()))
        concentration_prefix = re.search(rf"(?P<q>{NUMBER}\s*(?:mol\s*/\s*L|mM|M))\s+(?:aqueous\s+)?$", normalized_text[:name_span[0]], re.I)
        if concentration_prefix:
            qm = QUANTITY_RE.search(concentration_prefix["q"])
            chemical.add_quantity(_parse_quantity(qm, concentration_prefix.start("q")))
        chemical.formulation |= bool(re.search(r"\d\s*%|\b(?:purity|dispersion)\b", body+" "+chemical.name, re.I))
        if re.search(rf"{NUMBER}\s*N\s*$",normalized_text[max(0,name_span[0]-30):name_span[0]],re.I):
            chemical.formulation = True
            chemical.warnings.append("Normality requires stoichiometric review; it is not silently converted to molarity.")
        if re.search(r"\d\s*[-–]\s*\d.*(?:equiv|mmol)", body):
            chemical.warnings.append("Loading range: choose one value before using the schema.")

    # Also support "2.0 equiv. (0.8 mmol) of Et3N·3HF".
    prefix_re = re.compile(
        rf"(?P<q1>{NUMBER}\s*{UNIT})(?:\s*\((?P<q2>[^()]{{1,140}})\))?\s+(?:of\s+)?",
        re.IGNORECASE,
    )
    for match in prefix_re.finditer(normalized_text):
        if re.search(r"\byield\s*:\s*$", normalized_text[max(0,match.start()-20):match.start()],re.I):
            continue
        first_quantity = QUANTITY_RE.search(match["q1"])
        if _quantity_kind(first_quantity["unit"]) in {"percentage","concentration_m"}:
            continue
        if any(a <= match.start() < b for a,b in consumed_quantities):
            continue
        if any(a < match.start() < b for a,_,b in loading_blocks):
            continue
        name, name_span = _name_right_of(normalized_text, match.end())
        formulation_prefix = re.match(rf"(?:aqueous\s+)?(?P<q>{NUMBER}\s*(?:mol\s*/\s*L|mM|M|wt\s*%|%))\s+(?:aqueous\s+)?", name, re.I)
        formulation_quantity = None
        if formulation_prefix:
            formulation_quantity = QUANTITY_RE.search(formulation_prefix["q"])
            formulation_offset = name_span[0] + formulation_prefix.start("q")
            name_span = (name_span[0]+formulation_prefix.end(), name_span[1])
            name = name[formulation_prefix.end():]
        if "(" in name and QUANTITY_RE.search(name) or re.search(r"\b(?:mixture|solution|portion)\b",name,re.I):
            continue
        chemical = chemical_for(name, name_span, context=normalized_text[max(0, match.start() - 40) : match.start()])
        if chemical is None:
            continue
        if formulation_quantity:
            chemical.add_quantity(_parse_quantity(formulation_quantity, formulation_offset))
            chemical.formulation |= _quantity_kind(formulation_quantity["unit"]) == "percentage"
        for group in ("q1", "q2"):
            raw = match.group(group)
            if not raw:
                continue
            for quantity_match in QUANTITY_RE.finditer(raw):
                chemical.add_quantity(_parse_quantity(quantity_match, match.start(group)))

    # User-marked variable names are useful even when the method omits loading.
    for requested in variable_names:
        if any(_same_chemical_name(requested, c.name) for c in chemicals_by_name.values()):
            continue
        needle = _normalize_text(requested)
        mention = re.search(r"(?<![A-Za-z0-9])" + re.escape(needle) + r"(?![A-Za-z0-9])", normalized_text, re.I)
        if mention:
            chemical_for(mention[0], mention.span())

    # Symbolic optimization loadings still identify a material; leave the
    # unresolved loading literal rather than inventing a numeric value.
    for symbolic in re.finditer(r"\(\s*[XYZxyz]\s*(?:equiv\.?|eq\.?|mol\s*%)\s*\)", normalized_text):
        name, span, context = _name_left_of(normalized_text, symbolic.start())
        chemical_for(name, span, context=context)

    # Explicitly named reaction solvents can be unquantified; keep them visible
    # and mark their concentration missing instead of silently dropping them.
    workup_start = WORKUP_CONTEXT_RE.search(normalized_text)
    reaction_prefix = normalized_text[:workup_start.start()] if workup_start else normalized_text
    reaction_prefix = re.split(r"\b(?:TLC|concentrated|evaporation|removed|poured)\b",reaction_prefix,flags=re.I)[0]
    for solvent in sorted(SOLVENT_NAMES, key=len, reverse=True):
        if solvent in {"solvent","ether","hexane","n-hexane"}: continue
        for mention in re.finditer(r"(?<![A-Za-z0-9])"+re.escape(solvent)+r"(?![A-Za-z0-9])",reaction_prefix,re.I):
            if re.search(r"\b(?:bath|coolant|Dewar)\b",reaction_prefix[max(0,mention.start()-35):mention.end()+20],re.I): continue
            if any(a < mention.start() < b for a,_,b in loading_blocks): continue
            if any(a <= mention.start() < b for c in chemicals_by_name.values() for a,b in c.name_spans): continue
            if any(normalize_reagent_name(c.name) == normalize_reagent_name(mention[0]) for c in chemicals_by_name.values()): continue
            chemical_for(mention[0],mention.span(),context="solvent ")

    chemicals = list(chemicals_by_name.values())
    for chemical in chemicals:
        if (chemical.role == "reagent" and _normalize_name(chemical.name) == "glacial acetic acid"
                and chemical.value("volume_ml") is not None
                and chemical.value("amount_mmol") is None and chemical.value("equivalents") is None):
            chemical.role = "solvent"
    _apply_builtin_catalog(chemicals)
    inventory = _read_inventory(inventory_path)
    _apply_inventory(chemicals, inventory)
    _derive_values(chemicals, None)
    reference = next((c for name in variable_names[:1] for c in chemicals if _same_chemical_name(name,c.name) and c.role != "workup"), None) or _pick_reference_chemical(chemicals)
    if reference is not None and reference.value("equivalents") is None:
        _add_derived(reference,"equivalents",1.0,"equiv","selected reference reagent")
    reference_mmol = _reference_amount_mmol(reference)
    _derive_values(chemicals, reference_mmol)
    _assign_variables(chemicals, variable_names, reference)
    _assign_aliases(chemicals)

    unresolved: list[dict[str, str]] = []
    if reference is not None and not variable_names:
        unresolved.append({"chemical":reference.name,"alias":reference.alias,"field":"variable_assignment",
                           "issue":"Reference and variable substrate were inferred from one example; confirm the compounds that vary across Scope."})
    for chemical in chemicals:
        if chemical.role == "workup":
            continue
        for kind in ("mass_mg","amount_mmol","volume_ml"):
            explicit_values = [q.value for q in chemical.quantities if q.kind == kind and not q.derived]
            if len(set(explicit_values)) > 1:
                chemical.warnings.append(f"Repeated {kind} additions differ; stage/allocation review required.")
        if chemical.source_quantity("equivalents") and chemical.value("amount_mmol") is not None and reference_mmol:
            expected_eq = chemical.value("amount_mmol")/reference_mmol
            if abs(expected_eq-chemical.value("equivalents")) > max(.01,abs(expected_eq)*.03):
                chemical.warnings.append(f"Equivalents conflict: stated {chemical.value('equivalents'):.5g}, amount/reference gives {expected_eq:.5g}.")
        if chemical.variable:
            unresolved.append({"chemical":chemical.name,"alias":chemical.alias,"field":"structure",
                               "issue":"Insert the structure for each Scope row; its name and MW are compound-specific."})
        if chemical.role != "solvent" and not chemical.variable:
            if chemical.value("equivalents") is not None and chemical.molecular_weight_g_mol is None:
                unresolved.append(
                    {
                        "chemical": chemical.name,
                        "alias": chemical.alias, "field": "molecular_weight_g_mol",
                        "issue": "MW is required to calculate mass; add it to the generated Reaction_schema.docx.",
                    }
                )
        if chemical.role == "solvent" and chemical.value("concentration_m") is None:
            unresolved.append(
                {
                    "chemical": chemical.name,
                    "alias": chemical.alias, "field": "concentration_m",
                    "issue": "Reaction concentration could not be inferred; fill Concentration, M manually.",
                }
            )
        if chemical.role != "solvent" and chemical.value("equivalents") is None:
            unresolved.append({"chemical":chemical.name,"alias":chemical.alias,"field":"equivalents","issue":"Specify equivalents or a loading relative to Reagent_1."})
        if chemical.role not in {"solvent","workup"} and chemical.value("volume_ml") is not None and chemical.value("concentration_m") is None and chemical.density_g_ml is None:
            unresolved.append({"chemical":chemical.name,"alias":chemical.alias,"field":"density_g_ml","issue":"Volume calculation requires density of the actual material or stock-solution concentration."})
        if chemical.formulation:
            unresolved.append({"chemical":chemical.name,"alias":chemical.alias,"field":"formulation","issue":"Confirm purity/support/solution formulation before mass calculation."})
        for warning in dict.fromkeys(chemical.warnings):
            if any(term in warning for term in ("conflict", "Repeated", "range", "Normality", "stage review")):
                unresolved.append({"chemical":chemical.name,"alias":chemical.alias,"field":"review","issue":warning})

    if not chemicals:
        unresolved.append(
            {
                "chemical": "",
                "issue": "No explicit chemical loading was detected. Mark reagents with quantities or review manually.",
            }
        )

    if re.search(r"\beither\b.*\bor\b|\b(?:[XYZxyz] mol%|[XYZxyz] equiv|crude.*(?:next step|redissolved|dissolved))", normalized_text, re.I):
        unresolved.append({"chemical":"","field":"method","issue":"Conditional or multistage procedure: select conditions and split reaction stages before using one schema."})
    first_workup = WORKUP_CONTEXT_RE.search(normalized_text)
    if first_workup and any(c.role != "workup" and any(s[0]>first_workup.start() for s in c.name_spans) for c in chemicals):
        unresolved.append({"chemical":"","field":"method","issue":"Reaction additions occur after a filtration/workup marker; confirm reagent preparation or separate reaction stages before scaling."})

    result = {
        "schema_version": "1.1",
        "source_text": normalized_text,
        "reference_chemical": reference.name if reference else None,
        "reference_amount_mmol": reference_mmol,
        "chemicals": [_chemical_to_dict(chemical) for chemical in chemicals],
        "template_text": _build_template_text(normalized_text, chemicals),
        "ignored_workup_entities": [chemical.name for chemical in chemicals if chemical.role == "workup"],
        "unresolved": unresolved,
        "requires_user_confirmation": True,
        "catalog": catalog_summary(),
        "scope_requirement": (
            "Insert editable ChemDraw OLE structures, product numbers, limiting-reagent masses, "
            "and isolated product masses into Scope.docx."
        ),
    }
    return result


def generate_procedure_inputs(
    text: str,
    output_folder: str | Path,
    *,
    variable_names: Iterable[str] = (),
    inventory_path: str | Path | None = None,
    product_numbers: Iterable[str] = (),
    source_method: str | Path | None = None,
) -> GeneratedProcedureInputs:
    """Generate a complete editable input package from one ordinary procedure."""
    folder = Path(output_folder)
    folder.mkdir(parents=True, exist_ok=True)
    parsed = parse_procedure(text, variable_names=variable_names, inventory_path=inventory_path)

    compound_table_path = folder / "Compound_table.docx"
    schema_path = folder / "Reaction_schema.docx"
    scope_path = folder / "Scope.docx"
    template_path = folder / "SI_template.docx"
    all_in_one_path = folder / "All_in_one_input.docx"
    report_path = folder / "procedure_import_report.json"
    loadings_table_path = folder / "Loadings_table.docx"
    source_method_path: Path | None = None

    _write_compound_table(compound_table_path, product_numbers)
    _write_reaction_schema(parsed, schema_path)
    _write_scope_draft(parsed, scope_path, product_numbers)
    _write_si_template(parsed, template_path)
    _write_loadings_table(parsed, loadings_table_path)
    build_unified_input_docx(
        compound_table_path,
        all_in_one_path,
        reaction_schema=schema_path,
        scope=scope_path,
        si_template=template_path,
    )
    if source_method:
        source = Path(source_method).expanduser().resolve()
        if not source.exists() or not source.is_file():
            raise FileNotFoundError(f"Procedure document does not exist: {source}")
        source_method_path = folder / "Method.docx"
        shutil.copy2(source, source_method_path)
    report_path.write_text(json.dumps(parsed, ensure_ascii=False, indent=2), encoding="utf-8")
    return GeneratedProcedureInputs(
        compound_table=compound_table_path,
        reaction_schema=schema_path,
        scope=scope_path,
        si_template=template_path,
        all_in_one=all_in_one_path,
        report=report_path,
        loadings_table=loadings_table_path,
        source_method=source_method_path,
    )


def generate_procedure_inputs_from_docx(
    procedure_docx: str | Path,
    output_folder: str | Path,
    *,
    variable_names: Iterable[str] = (),
    inventory_path: str | Path | None = None,
    product_numbers: Iterable[str] = (),
) -> GeneratedProcedureInputs:
    """Read a Word method and generate both classic and all-in-one inputs."""
    source = Path(procedure_docx).expanduser().resolve()
    if source.suffix.casefold() != ".docx":
        raise ValueError("Method must be a .docx document.")
    return generate_procedure_inputs(
        read_procedure_text(source),
        output_folder,
        variable_names=variable_names,
        inventory_path=inventory_path,
        product_numbers=product_numbers,
        source_method=source,
    )


def _write_compound_table(path: Path, product_numbers: Iterable[str]) -> None:
    document = Document()
    _set_default_font(document)
    _add_document_title(document, "Compound table")
    document.add_paragraph(
        "One row is one product. Insert an editable ChemDraw OLE structure, fill measured values, "
        "and enter - for an optional block that must be omitted from the final SI. Product numbers "
        "must match Scope.docx and spectra folder names."
    )
    headers = ("number", "structure", "color", "mp", "Rf", "HRMS", "Elemental_analysis", "IR")
    numbers = [number.strip() for number in product_numbers if number.strip()] or [""]
    table = document.add_table(rows=1 + len(numbers), cols=len(headers))
    table.style = "Table Grid"
    for column, heading in enumerate(headers):
        table.cell(0, column).text = heading
        for run in table.cell(0, column).paragraphs[0].runs:
            run.bold = True
    for row_index, number in enumerate(numbers, start=1):
        table.cell(row_index, 0).text = number
        for cell in table.rows[row_index].cells:
            if not cell.text.strip():
                _shade_cell(cell)
    document.save(path)


def _write_reaction_schema(parsed: dict[str, Any], path: Path) -> None:
    document = Document()
    _set_default_font(document)
    table = document.add_table(rows=1, cols=len(SCHEMA_HEADERS))
    table.style = "Table Grid"
    for index, heading in enumerate(SCHEMA_HEADERS):
        table.cell(0, index).text = heading
        for run in table.cell(0, index).paragraphs[0].runs:
            run.bold = True

    for chemical in parsed["chemicals"]:
        if chemical["role"] == "workup":
            continue
        row = table.add_row().cells
        row[0].text = chemical["alias"]
        row[1].text = _format_number(chemical["equivalents"])
        row[2].text = "" if chemical["variable"] or chemical.get("formulation") else _format_number(chemical["molecular_weight_g_mol"])
        row[3].text = _format_number(chemical["density_g_ml"])
        row[4].text = _format_number(chemical["concentration_m"])
        fields = {issue.get("field") for issue in parsed["unresolved"] if issue.get("alias") == chemical["alias"]}
        for field,column in (("equivalents",1),("molecular_weight_g_mol",2),("density_g_ml",3),("concentration_m",4)):
            if field in fields: _shade_cell(row[column])
        if chemical["variable"] or chemical.get("formulation"): _shade_cell(row[2])
        if fields & {"review","formulation"}: _shade_cell(row[0])
    document.add_paragraph("Yellow cells require your input or review. Variable-reagent MW comes from its Scope structure. Blank solvent MW/density cells are not necessarily required.")
    catalog = parsed.get("catalog",{})
    document.add_paragraph(f"Offline catalog: {catalog.get('compound_count',0):,} compounds. Property sources and reference temperatures are saved in procedure_import_report.json.")
    for chemical in parsed["chemicals"]:
        if chemical["role"] != "workup":
            document.add_paragraph(f"{chemical['alias']}: {chemical['name']}")
    document.save(path)


def _write_scope_draft(parsed: dict[str, Any], path: Path, product_numbers: Iterable[str]) -> None:
    variables = [chemical for chemical in parsed["chemicals"] if chemical["variable"]]
    document = Document()
    _set_default_font(document)
    _add_document_title(document, "Scope from procedure")
    document.add_paragraph(
        "Replace the blank structure cells with editable ChemDraw OLE objects. Fill the actual mass of "
        "Reagent_1 and the isolated product mass for every product. Product_number must match Compound_table.docx."
    )
    headers: list[str] = []
    for chemical in variables:
        headers.append(chemical["alias"])
        if chemical["alias"] == "Reagent_1":
            headers.append("Mass of Reagent_1, mg")
    if not variables:
        headers.extend(("Reagent_1", "Mass of Reagent_1, mg"))
    headers.extend(("Product", "Product_number", "Mass of product, mg"))
    numbers = [number.strip() for number in product_numbers if number.strip()] or [""]
    table = document.add_table(rows=1 + len(numbers), cols=len(headers))
    table.style = "Table Grid"
    for column, heading in enumerate(headers):
        table.cell(0, column).text = heading
        for run in table.cell(0, column).paragraphs[0].runs:
            run.bold = True
    product_number_column = len(headers) - 2
    for row_index, number in enumerate(numbers, start=1):
        table.cell(row_index, product_number_column).text = number
        for cell in table.rows[row_index].cells:
            if not cell.text.strip(): _shade_cell(cell)
    document.save(path)


def _write_si_template(parsed: dict[str, Any], path: Path) -> None:
    base_template = Path(__file__).resolve().parent / "templates" / "SI_template.docx"
    document = Document(base_template)
    paragraphs = document.paragraphs
    if len(paragraphs) < 5:
        raise ValueError(f"Bundled SI template is incomplete: {base_template}")
    paragraph = paragraphs[2]
    paragraph.clear()
    for obsolete in paragraphs[3:5]:
        obsolete._element.getparent().remove(obsolete._element)
    issues_by_alias: dict[str, set[str]] = {}
    for issue in parsed["unresolved"]:
        issues_by_alias.setdefault(issue.get("alias",""),set()).add(issue.get("field","review"))
    unresolved_names = sorted({c["name"] for c in parsed["chemicals"] if not c["variable"] and issues_by_alias.get(c["alias"])}, key=len, reverse=True)
    name_pattern = re.compile(r"(?<!\w)(?:" + "|".join(re.escape(n) for n in unresolved_names) + r")(?!\w)") if unresolved_names else None

    def add_literal(text: str) -> None:
        position = 0
        for name in name_pattern.finditer(text) if name_pattern else []:
            paragraph.add_run(text[position:name.start()])
            paragraph.add_run(name[0]).font.highlight_color = WD_COLOR_INDEX.YELLOW
            position = name.end()
        paragraph.add_run(text[position:])

    template_text = parsed["template_text"].strip()
    if template_text and template_text[-1] not in ".!?":
        template_text += "."
    template_text += (
        " Yield {Product.mg} mg ({Product.yield.percent}); {Product.appearance}; "
        "mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system})."
    )
    offset = 0
    for match in re.finditer(r"\{(?P<alias>[^{}.]+)\.(?P<attr>[^{}]+)\}", template_text):
        add_literal(template_text[offset:match.start()])
        run = paragraph.add_run(match[0])
        if issues_by_alias.get(match["alias"]):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        offset = match.end()
    add_literal(template_text[offset:])
    document.save(path)


def _write_loadings_table(parsed: dict[str, Any], path: Path) -> None:
    """Write a human-review table of source and normalized reaction loadings."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Inches(0.35)
    section.top_margin = section.bottom_margin = Inches(0.45)
    _set_default_font(document)
    _add_document_title(document, "Compound loadings extracted from procedure")
    document.add_paragraph(
        "Values are normalized to mg, mmol, mL, equivalents and mol/L. Yellow cells require manual input "
        "or review. N/A means the property is not normally required for that material. Calculated "
        "values are explicitly identified in the last column."
    )
    table = document.add_table(rows=1, cols=len(LOADINGS_HEADERS))
    table.style = "Table Grid"
    table.autofit = False
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for index, heading in enumerate(LOADINGS_HEADERS):
        cell = table.cell(0, index)
        cell.text = heading
        cell.width = Inches(LOADINGS_COLUMN_WIDTHS_INCHES[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cell, "1F4E78")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    issues_by_alias: dict[str, list[dict[str, str]]] = {}
    for issue in parsed["unresolved"]:
        if issue.get("alias"):
            issues_by_alias.setdefault(issue["alias"], []).append(issue)
    field_columns = {
        "mass_mg": 4,
        "amount_mmol": 5,
        "volume_ml": 6,
        "equivalents": 7,
        "concentration_m": 8,
        "molecular_weight_g_mol": 9,
        "density_g_ml": 10,
    }
    quantity_names = {
        "mass_mg": "mass",
        "amount_mmol": "amount",
        "volume_ml": "volume",
        "equivalents": "equivalents",
        "concentration_m": "concentration",
        "percentage": "percentage",
    }

    for chemical in parsed["chemicals"]:
        if chemical["role"] == "workup":
            continue
        row = table.add_row().cells
        for index, cell in enumerate(row):
            cell.width = Inches(LOADINGS_COLUMN_WIDTHS_INCHES[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row[0].text = chemical["alias"]
        row[1].text = chemical["name"]
        row[2].text = chemical["role"] + ("; variable" if chemical["variable"] else "; constant")
        explicit = [q for q in chemical["quantities"] if not q["derived"] and q.get("raw")]
        row[3].text = "; ".join(dict.fromkeys(q["raw"] for q in explicit)) or "not stated"
        for key,column in field_columns.items():
            row[column].text = _format_number(chemical.get(key)) or ""

        derived = [q for q in chemical["quantities"] if q["derived"]]
        review = [issue["issue"] for issue in issues_by_alias.get(chemical["alias"], [])]
        provenance: list[str] = []
        explicit_kinds = list(dict.fromkeys(quantity_names.get(q["kind"],q["kind"]) for q in explicit))
        if explicit_kinds:
            provenance.append("text: " + ", ".join(explicit_kinds))
        for quantity in derived:
            provenance.append(
                "calculated " + quantity_names.get(quantity["kind"],quantity["kind"])
                + (f" ({quantity['formula']})" if quantity.get("formula") else "")
            )
        for prop,label in (("molecular_weight_g_mol","MW"),("density_g_ml","density")):
            source = chemical.get("property_sources",{}).get(prop)
            if source:
                provenance.append(f"{label}: {source.get('source','reference')}")
        if review:
            provenance.append("REVIEW: " + " | ".join(dict.fromkeys(review)))
        row[11].text = "; ".join(provenance) or "No loading source found"

        applicable = {
            "mass_mg": chemical["role"] != "solvent",
            "amount_mmol": chemical["role"] != "solvent",
            "volume_ml": chemical["role"] == "solvent" or chemical.get("volume_ml") is not None,
            "equivalents": chemical["role"] != "solvent",
            "concentration_m": chemical["role"] == "solvent" or chemical.get("concentration_m") is not None,
            "molecular_weight_g_mol": chemical["role"] != "solvent",
            "density_g_ml": (
                chemical.get("density_g_ml") is not None
                or chemical["role"] != "solvent" and chemical.get("volume_ml") is not None
                and chemical.get("concentration_m") is None
            ),
        }
        unresolved_fields = {issue.get("field") for issue in issues_by_alias.get(chemical["alias"], [])}
        for key,column in field_columns.items():
            if not applicable[key]:
                row[column].text = "N/A"
            elif chemical.get(key) is None or key in unresolved_fields:
                _shade_cell(row[column])
        if review or chemical.get("warnings") or chemical.get("formulation"):
            _shade_cell(row[11])
        if chemical["variable"]:
            _shade_cell(row[1])

    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                if column_index in field_columns.values():
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(7.25)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    if parsed.get("ignored_workup_entities"):
        document.add_paragraph(
            "Work-up materials excluded from the reaction-loading table: "
            + ", ".join(parsed["ignored_workup_entities"])
            + ". Their quantities remain literal in SI_template.docx."
        )
    document.add_paragraph(
        "Reaction_schema.docx remains the machine-readable calculation schema. This table is an auditable "
        "summary of the example procedure and must be reviewed before scaling."
    )
    document.save(path)


def _shade_cell(cell, fill: str = "FFF2CC") -> None:
    properties = cell._tc.get_or_add_tcPr()
    for existing in properties.findall(qn("w:shd")):
        properties.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)


def _add_document_title(document: Document, text: str) -> None:
    style = document.styles["Title"]
    style.font.color.rgb = RGBColor(0, 0, 0)
    style_properties = style.element.get_or_add_pPr()
    border = style_properties.find(qn("w:pBdr"))
    if border is not None:
        style_properties.remove(border)
    paragraph = document.add_paragraph(text, style="Title")
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


_CHEMISTRY_KINDS = {
    "amount_mmol",
    "concentration_m",
    "equivalents",
    "mass_mg",
    "percentage",
    "volume_ml",
}


def _parse_quantity(match: re.Match[str], offset: int) -> ParsedQuantity:
    numeric = match.group("value").replace(",", ".")
    if "/" in numeric:
        numerator, denominator = numeric.split("/")
        value = float(numerator)/float(denominator)
    else:
        value = float(numeric)
    source_unit = re.sub(r"\s+", "", match.group("unit"))
    key = source_unit.casefold().rstrip(".")
    kind = _quantity_kind(source_unit)
    canonical_unit = source_unit
    if key in {"µl", "μl", "ul"}:
        value, canonical_unit = value / 1000, "mL"
    elif key == "l":
        value, canonical_unit = value * 1000, "mL"
    elif key in {"µmol", "μmol", "umol"}:
        value, canonical_unit = value / 1000, "mmol"
    elif key == "mol":
        value, canonical_unit = value * 1000, "mmol"
    elif key in {"µg", "μg", "ug"}:
        value, canonical_unit = value / 1000, "mg"
    elif key == "g":
        value, canonical_unit = value * 1000, "mg"
    elif key == "kg":
        value, canonical_unit = value * 1_000_000, "mg"
    elif key == "mm":
        value, canonical_unit = value / 1000, "M"
    elif kind == "equivalents":
        canonical_unit = "equiv"
    elif kind == "percentage":
        canonical_unit = source_unit
    return ParsedQuantity(
        kind=kind,
        value=value,
        canonical_unit=canonical_unit,
        source_unit=source_unit,
        raw=match.group(0),
        span=(offset + match.start(), offset + match.end()),
    )


def _quantity_kind(unit: str) -> str:
    key = re.sub(r"\s+", "", unit).casefold().rstrip(".")
    if key in {"µl", "μl", "ul", "ml", "l"}:
        return "volume_ml"
    if key in {"µmol", "μmol", "umol", "mmol", "mol"}:
        return "amount_mmol"
    if key in {"µg", "μg", "ug", "mg", "g", "kg"}:
        return "mass_mg"
    if key in {"equiv", "equivalent", "equivalents", "eq"}:
        return "equivalents"
    if key in {"m", "mm", "mol/l"}:
        return "concentration_m"
    if key in {"mol%", "wt%", "%"}:
        return "percentage"
    return "unknown"


def _name_left_of(text: str, position: int) -> tuple[str, tuple[int, int], str]:
    while position and text[position-1] in " ,": position -= 1
    # Only delimiters outside chemical/quantity parentheses can split a name.
    depth = 0
    start = 0
    visible = list(text[:position])
    for index, char in enumerate(text[:position]):
        if char == "(": depth += 1
        if depth: visible[index] = " "
        if char == ")": depth = max(0,depth-1)
        locant_comma = char=="," and index>0 and text[index-1].isdigit() and re.match(r"\s*\d",text[index+1:position])
        if not depth and (char in ";\n:" or char in ",." and index+1 < position and text[index+1].isspace() and not locant_comma):
            start = index+1
    segment = text[start:position]
    visible_segment = "".join(visible[start:position])
    connectors = list(re.finditer(
        r"\b(?:solution|suspension|mixture)\s+of\b|\bin(?=\d+(?:,\d+)+-)|\b(?:containing|added|add|loaded|charged|with|and|or|either|in|of|by|then|thereafter|after which|after that|before)\b|(?<=with)(?=[A-Za-z])",
        visible_segment,re.I))
    relative_start = connectors[-1].end() if connectors else 0
    raw = segment[relative_start:]
    # Concentration belongs to the stock, not to its identity.
    prefix = re.match(rf"\s*(?:{NUMBER}\s*(?:mol\s*/\s*L|mM|M|N)\s+)?(?:aqueous\s+)?",raw,re.I)
    if prefix:
        relative_start += prefix.end()
        raw = segment[relative_start:]
    leading = len(raw) - len(raw.lstrip())
    trailing = len(raw.rstrip())
    name_start = start + relative_start + leading
    name_end = start + relative_start + trailing
    context = text[max(0, name_start - 180) : name_start]
    return text[name_start:name_end], (name_start, name_end), context


def _name_right_of(text: str, position: int) -> tuple[str, tuple[int, int]]:
    tail = text[position : position + 260]
    # A following loading block is not part of the name. Structural and label
    # parentheses remain intact, including commas inside a structural name.
    loading_starts = [a for a,_,_ in _loading_parentheses(tail)]
    if loading_starts: tail = tail[:min(loading_starts)]
    visible = list(tail)
    depth = 0
    for i,char in enumerate(tail):
        if char == "(": depth += 1
        if depth: visible[i] = " "
        if char == ")": depth = max(0,depth-1)
    visible_text = "".join(visible)
    visible_text = re.sub(r"(?<=\d),(?=\s*\d)", " ", visible_text)
    visible_text = re.sub(r"\banhyd\.", lambda m:m[0][:-1]+" ", visible_text, flags=re.I)
    stop = re.search(
        r",\s|;|\.\s|\band\b|\bin\b|\bat\b|\bunder\b|\bfor\b|\bto\b|\bbefore\b|\bafter\b|\bfollowed\b|\bwas\b|\bwere\b|\bis\b|\bare\b",
        visible_text,
        re.IGNORECASE,
    )
    raw = tail[: stop.start()] if stop else tail
    leading = len(raw) - len(raw.lstrip())
    cleaned = raw.strip(" .,:;")
    start = position + leading
    return cleaned, (start, start + len(cleaned))


def _clean_name(name: str) -> str:
    value = re.sub(r"\s+", " ", name).strip(" .,:;-/")
    value = re.sub(r"^(?:a|an|the)\s+", "", value, flags=re.IGNORECASE)
    value = re.sub(r"^(?:after\s+that|dropwise|then|subsequently)\s+", "", value, flags=re.I)
    value = re.sub(r"^To\s+a\s+(.+?)\s+solution$",r"\1",value,flags=re.I)
    value = re.sub(r"^To\s+a\s+s\s+u\s+s\s+p\s+e\s+n\s+s\s+i\s+o\s+n\s+o\s+f\s+", "",value,flags=re.I)
    value = re.sub(r"^solid\s+", "",value,flags=re.I)
    return value


def _plausible_chemical_name(name: str) -> bool:
    if not name or len(name) > 220:
        return False
    if QUANTITY_RE.fullmatch(name): return False
    if re.fullmatch(r"\d+[a-z]?",name,re.I):
        return True
    if not re.search(r"[A-Za-zА-Яа-я]",name): return False
    lowered = _normalize_name(name)
    if re.search(r"\b(?:plate|electrode|mesh|lamp|leds?|reactor|cuvette|syringe|funnel|diameter|aliquot|giving|affording|yielding|furnishing)\b", lowered):
        return False
    if lowered in {"organic","aqueous","organic phase","aqueous phase","ice","ice-water"}: return False
    return not any(
        token in lowered
        for token in ("flask", "vial", "stirring bar", "stir bar", "condenser", "reaction mixture", "oil bath",
                      "afford", "gave ", "give ", "delivered", "furnish", "obtain", "purification", "yield",
                      "then ", "was ", "were ", "added", "stirred", "equiv", "mmol", "atmosphere", "tube")
    )


def _normalize_text(text: str) -> str:
    value = (
        text.replace("\u00a0", " ")
        .replace("μ", "µ")
        .replace("−", "-")
        .replace("–", "-")
        .replace("\uf06d", "µ")
        .translate(str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789"))
        .strip()
    )
    return re.sub(r"\s+", " ", value)


def _normalize_name(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().casefold())


def _infer_role(name: str, context: str) -> str:
    normalized = _normalize_name(name)
    normalized = re.sub(r"^(?:(?:dry|anhydrous|anhyd\.?|degassed|freshly|distilled|deionized)\s+)+", "", normalized)
    # Workup state is decided by the full-text state machine in chemical_for,
    # not a truncated context that may still contain the previous stage.
    local_context = re.split(r"[.;]\s+(?=[A-Z])",context)[-1]
    solvent_keys = {normalize_reagent_name(n) for n in SOLVENT_NAMES}
    parenthetic_alias = re.search(r"\(([^()]+)\)$",normalized)
    if normalize_reagent_name(normalized) in solvent_keys or parenthetic_alias and normalize_reagent_name(parenthetic_alias[1]) in solvent_keys:
        return "solvent"
    if _solvent_mixture_parts(name):
        return "solvent"
    if re.search(r"\b(?:solvent|dissolved\s+in|suspended\s+in)\s*$", local_context, re.IGNORECASE):
        return "solvent"
    return "reagent"


def _solvent_mixture_parts(name: str) -> list[str]:
    parts = re.split(r"\s+and\s+|\s*[/,:]\s*", name, flags=re.I)
    known = {normalize_reagent_name(n) for n in SOLVENT_NAMES}
    return parts if len(parts) > 1 and all(normalize_reagent_name(p) in known for p in parts) else []


def _mixture_left_of(text: str, position: int) -> tuple[str, tuple[int, int]] | None:
    start = max(0, position-150)
    prefix = text[start:position].rstrip()
    for match in reversed(list(re.finditer(r"\b(?:mixture\s+of|in)\s+", prefix, re.I))):
        names = prefix[match.end():].strip()
        if _solvent_mixture_parts(names):
            return names, (start+match.end(), start+match.end()+len(names))
    suffix = re.search(r"[A-Za-z0-9,-]+(?:\s*(?:[:/]|\band\b)\s*[A-Za-z0-9,-]+)+$",prefix,re.I)
    if suffix and _solvent_mixture_parts(suffix[0]):
        return suffix[0], (start+suffix.start(),start+suffix.end())
    return None


def _stock_left_of(text: str, position: int, quantities) -> tuple[str, tuple[int,int], str] | None:
    kinds={_quantity_kind(q["unit"]) for q in quantities}
    prefix=text[:position].rstrip()
    carrier=re.search(r"\bin\s+(?P<solvent>[A-Za-z0-9,-]+)$",prefix,re.I)
    known={normalize_reagent_name(n) for n in SOLVENT_NAMES}
    if not carrier or normalize_reagent_name(carrier["solvent"]) not in known: return None
    name,span,_=_name_left_of(text,carrier.start())
    if not _plausible_chemical_name(name): return None
    concentration_before = re.search(rf"{NUMBER}\s*(?:mol\s*/\s*L|mM|M)\s*$",text[:span[0]],re.I)
    if ("concentration_m" not in kinds and not concentration_before) or not kinds.intersection({"amount_mmol","equivalents","volume_ml"}):
        return None
    return text[span[0]:len(prefix)], (span[0],len(prefix)), carrier["solvent"]


def _is_workup_at(text: str, position: int) -> bool:
    """Carry workup state across sentences, resetting only at an explicit stage."""
    prefix = text[:position]
    events = [(m.start(), True) for m in WORKUP_CONTEXT_RE.finditer(prefix)]
    events.extend((m.start(),True) for m in re.finditer(r"\bpoured\s+into\s+(?:(?:a|the)\s+)?(?:mixture\s+of\s+)?(?:\d+(?:\.\d+)?\s*mL\s+of\s+)?ice\b",prefix,re.I))
    restarts = re.finditer(r"\b(?:step|stage)\s*[2-9]\b|\bTo\s+(?:a|the)\s+(?:(?:stirred|heterogeneous|cooled)\s+)*solution\b|\bnext\s+(?:reaction|step)\b[^.]*\.",prefix,re.I)
    events.extend((m.end(), False) for m in restarts)
    vessel_starts = re.finditer(r"\b(?:To|Into|In)\s+(?:a|an|the)\s+[^.;]{0,100}?\b(?:flask|vial|tube)\b|\bUnder\s+(?:a|an)?\s*(?:nitrogen|argon)\s+atmosphere\b",prefix,re.I)
    events.extend((m.end(), False) for m in vessel_starts)
    if not events: return False
    last_position,workup = sorted(events)[-1]
    if workup and re.search(r"\b(?:crude|residue|intermediate)\b",prefix[last_position:],re.I):
        following = text[position:position+450]
        # Do not mistake decimal loading points (e.g. 1.1 equiv.) for a
        # sentence boundary while recognising a subsequent transformation.
        next_workup = WORKUP_CONTEXT_RE.search(following)
        reaction_tail = following[:next_workup.start()] if next_workup else following
        if re.search(r"\b(?:added|treated)\b[\s\S]*\bstirred\b",reaction_tail,re.I):
            return False
    return workup


def _is_terminal_workup_addition(text: str, span: tuple[int,int]) -> bool:
    """Recognise a terminal quench/wash even when the author omits that verb."""
    following=text[span[1]:span[1]+260]
    preceding=text[max(0,span[0]-90):span[0]]
    residue_handling=bool(re.search(r"\b(?:residue|crude)\b[^.;]{0,70}\b(?:taken\s+up|dissolved|suspended)\s+in\s*$",preceding,re.I))
    if residue_handling:
        if re.search(r"\b(?:added|treated)\b[^.;]{0,180}\bstirred\b",following,re.I): return False
        if re.search(r"\b(?:filtered|filtration|purified|chromatography|recrystalli[sz]ed|triturated)\b",following,re.I): return True
    endpoint=re.search(r"\b(?:phases?\s+were\s+separated|organic\s+layer\s+was\s+separated|"
                       r"aqueous\s+layer|extracted\s+with|mixture\s+was\s+extracted|transferred\s+to\s+a\s+separatory)\b",following,re.I)
    if not endpoint: return False
    before_endpoint=following[:endpoint.start()]
    if re.search(r"\b(?:heated|refluxed|irradiated|stirred)\b[^.;]{0,50}\b(?:for|overnight)\b",before_endpoint,re.I):
        return False
    return bool(re.search(r"\b(?:aqueous|aq\.?|sat(?:urated)?\.?|HCl|NaOH|K2CO3|NaHCO3|NH4Cl|water|H2O|ice)\b",
                          text[max(0,span[0]-35):span[1]],re.I))


def _loading_parentheses(text: str):
    stack = []
    for index,char in enumerate(text):
        if char == "(": stack.append(index)
        elif char == ")" and stack:
            start = stack.pop()
            # Quantified loading parentheses may contain explanatory nested
            # structure names; innermost name parentheses have no quantities.
            body = text[start+1:index]
            if QUANTITY_RE.search(body) and not any(QUANTITY_RE.search(text[p+1:start]) for p in stack):
                yield start,start+1,index+1


def _same_chemical_name(left: str, right: str) -> bool:
    def stripped(value):
        value = _normalize_text(value).casefold()
        if " in " in value:
            active, carrier = value.rsplit(" in ",1)
            if normalize_reagent_name(carrier) in {normalize_reagent_name(n) for n in SOLVENT_NAMES}:
                value = active
        value = re.sub(r"^(?:(?:the|corresponding|appropriate|compounds?|substrates?|intermediates?|base|ligand|dry|anhydrous|anhyd\.?)\s+)+", "", value)
        return re.sub(r"\s+", "", value).strip(".,")
    return stripped(left) == stripped(right)


def _apply_builtin_catalog(chemicals: list[ParsedChemical]) -> None:
    for chemical in chemicals:
        if chemical.role == "workup": continue
        result = lookup_reagent(chemical.name)
        chemical.catalog_status = result["status"]
        record = result.get("record")
        if not record: continue
        chemical.molecular_weight_g_mol = record["molecular_weight_g_mol"]
        chemical.property_sources["molecular_weight_g_mol"] = {"source":"local_catalog","url":record["source_url"],"cid":record.get("cid"),"formula":record["formula"]}
        # Never apply neat density to a stock solution or impure formulation.
        if record.get("density_g_ml") is not None and chemical.value("concentration_m") is None and not chemical.formulation:
            chemical.density_g_ml = record["density_g_ml"]
            chemical.property_sources["density_g_ml"] = {"source":"local_catalog","url":record["density_source_url"],"temperature_c":record["density_temperature_c"],"basis":"neat liquid"}
        elif record.get("density_evidence"):
            chemical.property_sources["density_candidates"] = {"source":"local_pubchem_evidence", "approved_for_calculation":False,
                                                               "assertions":record["density_evidence"]}
        if chemical.value("mass_mg") is not None and chemical.value("amount_mmol"):
            observed = chemical.value("mass_mg")/chemical.value("amount_mmol")
            if abs(observed / chemical.molecular_weight_g_mol - 1) > .03:
                chemical.warnings.append(f"MW conflict: source mass/amount implies {observed:.5g} g/mol; catalog gives {chemical.molecular_weight_g_mol:.5g}. Check identity, purity, salt and hydrate.")


def _read_inventory(path: str | Path | None) -> dict[str, dict[str, str]]:
    if not path:
        return {}
    inventory: dict[str, dict[str, str]] = {}
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            names = [row.get("name", ""), *((row.get("synonyms") or "").split("|"))]
            for name in names:
                if name.strip():
                    inventory[_normalize_name(name)] = row
    return inventory


def _apply_inventory(chemicals: list[ParsedChemical], inventory: dict[str, dict[str, str]]) -> None:
    for chemical in chemicals:
        row = inventory.get(_normalize_name(chemical.name))
        if not row:
            continue
        mw = _float_or_none(
            row.get("molecular_weight_g_mol") or row.get("mw")
        )
        density = _float_or_none(row.get("density_g_ml") or row.get("density"))
        if mw is not None:
            chemical.molecular_weight_g_mol = mw
            chemical.property_sources["molecular_weight_g_mol"] = {"source":"user_inventory"}
        if density is not None:
            chemical.density_g_ml = density
            chemical.property_sources["density_g_ml"] = {"source":"user_inventory"}
        if row.get("role") and chemical.role != "workup":
            chemical.role = str(row["role"]).strip().casefold()


def _pick_reference_chemical(chemicals: list[ParsedChemical]) -> ParsedChemical | None:
    candidates = [chemical for chemical in chemicals if chemical.role not in {"solvent", "workup"}]
    if not candidates:
        return None
    return min(
        candidates,
        key=lambda chemical: (
            # Explicit substrate language is stronger evidence than list order.
            # This remains an assumption: a single example cannot define Scope.
            0 if re.search(r"\b(?:substrates?|starting material|coupling partner)\b", chemical.name, re.I) else 1,
            0 if chemical.value("amount_mmol") is not None else 1,
            abs((chemical.value("equivalents") or 1.0) - 1.0),
            0 if chemical.value("mass_mg") is not None else 1,
        ),
    )


def _reference_amount_mmol(reference: ParsedChemical | None) -> float | None:
    if reference is None:
        return None
    amount = reference.value("amount_mmol")
    equivalents = reference.value("equivalents") or 1.0
    return amount / equivalents if amount is not None and equivalents else None


def _derive_values(chemicals: list[ParsedChemical], reference_mmol: float | None) -> None:
    for chemical in chemicals:
        amount = chemical.value("amount_mmol")
        mass = chemical.value("mass_mg")
        volume = chemical.value("volume_ml")
        concentration = chemical.value("concentration_m")
        equivalents = chemical.value("equivalents")
        percentage = chemical.value("percentage")

        if chemical.molecular_weight_g_mol is None and mass is not None and amount and not chemical.formulation and chemical.catalog_status != "formulation":
            chemical.molecular_weight_g_mol = mass / amount
            chemical.property_sources["molecular_weight_g_mol"] = {"source":"source_mass_amount_ratio","requires_review":True}
            chemical.warnings.append("MW was inferred from the stated mass and amount; confirm it.")
        if chemical.density_g_ml is None and mass is not None and volume and concentration is None and not chemical.formulation:
            chemical.density_g_ml = mass / (volume * 1000)
            chemical.property_sources["density_g_ml"] = {"source":"source_mass_volume_ratio","requires_review":True}
            chemical.warnings.append("Density was inferred from the stated mass and volume; confirm it.")

        if chemical.role == "workup": continue
        if mass is None and volume is not None and chemical.density_g_ml is not None and concentration is None:
            _add_derived(
                chemical,
                "mass_mg",
                volume * chemical.density_g_ml * 1000,
                "mg",
                "volume_ml * density_g_ml * 1000",
            )
            mass = chemical.value("mass_mg")
        if amount is None and mass is not None and chemical.molecular_weight_g_mol and not chemical.formulation and chemical.role != "solvent":
            _add_derived(
                chemical,
                "amount_mmol",
                mass / chemical.molecular_weight_g_mol,
                "mmol",
                "mass_mg / molecular_weight_g_mol",
            )
            amount = chemical.value("amount_mmol")
        if amount is None and chemical.role != "solvent" and volume is not None and concentration is not None:
            _add_derived(
                chemical,
                "amount_mmol",
                volume * concentration,
                "mmol",
                "volume_ml * concentration_m",
            )
            amount = chemical.value("amount_mmol")
        if mass is None and amount is not None and chemical.molecular_weight_g_mol and not chemical.formulation:
            _add_derived(
                chemical,
                "mass_mg",
                amount * chemical.molecular_weight_g_mol,
                "mg",
                "amount_mmol * molecular_weight_g_mol",
            )

        percentage_quantity = chemical.source_quantity("percentage")
        if (
            equivalents is None
            and percentage is not None
            and percentage_quantity is not None
            and re.sub(r"\s+", "", percentage_quantity.source_unit).casefold() == "mol%"
        ):
            _add_derived(chemical, "equivalents", percentage / 100, "equiv", "mol_percent / 100")
            equivalents = chemical.value("equivalents")
        if equivalents is None and amount is not None and reference_mmol:
            _add_derived(chemical, "equivalents", amount / reference_mmol, "equiv", "amount / reference_amount")
            equivalents = chemical.value("equivalents")
        if chemical.role == "solvent" and chemical.value("concentration_m") is None and volume and reference_mmol:
            _add_derived(
                chemical,
                "concentration_m",
                reference_mmol / volume,
                "M",
                "reference_amount_mmol / solvent_volume_ml",
            )


def _add_derived(chemical: ParsedChemical, kind: str, value: float, unit: str, formula: str) -> None:
    if chemical.value(kind) is not None:
        return
    chemical.quantities.append(
        ParsedQuantity(kind, value, unit, unit, "", (-1, -1), derived=True, formula=formula)
    )


def _assign_variables(
    chemicals: list[ParsedChemical],
    variable_names: Iterable[str],
    reference: ParsedChemical | None,
) -> None:
    requested = [_normalize_name(name) for name in variable_names if name.strip()]
    if requested:
        matched: set[str] = set()
        variable_index = 1
        for wanted in requested:
            for chemical in chemicals:
                key = _normalize_name(chemical.name)
                if _same_chemical_name(wanted, key) and not chemical.variable and chemical.role != "workup":
                    chemical.variable = True
                    chemical.alias = f"Reagent_{variable_index}"
                    variable_index += 1
                    matched.add(wanted)
                    break
        missing = [name for name in requested if name not in matched]
        if missing and reference:
            reference.warnings.append("Requested variable names not found: " + ", ".join(missing))
        if not matched and reference is not None:
            reference.variable = True
            reference.alias = "Reagent_1"
    elif reference is not None:
        reference.variable = True
        reference.alias = "Reagent_1"


def _assign_aliases(chemicals: list[ParsedChemical]) -> None:
    variables = [chemical for chemical in chemicals if chemical.variable]
    used = {chemical.alias.casefold() for chemical in variables if chemical.alias}
    variable_index = 1
    for chemical in variables:
        if chemical.alias:
            continue
        while f"reagent_{variable_index}" in used:
            variable_index += 1
        chemical.alias = f"Reagent_{variable_index}"
        used.add(chemical.alias.casefold())
    for chemical in chemicals:
        if chemical.role == "workup":
            continue
        if chemical.variable:
            continue
        stem = _identifier(chemical.name)
        if chemical.role == "solvent":
            stem = f"Solvent_{stem}"
        alias = stem or "Reagent"
        base = alias
        suffix = 2
        while alias.casefold() in used:
            alias = f"{base}_{suffix}"
            suffix += 1
        chemical.alias = alias
        used.add(alias.casefold())


def _identifier(name: str) -> str:
    value = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")
    if value and value[0].isdigit():
        value = "R_" + value
    return value


def _chemical_to_dict(chemical: ParsedChemical) -> dict[str, Any]:
    return {
        "name": chemical.name,
        "alias": chemical.alias,
        "role": chemical.role,
        "variable": chemical.variable,
        "equivalents": chemical.value("equivalents"),
        "molecular_weight_g_mol": chemical.molecular_weight_g_mol,
        "density_g_ml": chemical.density_g_ml,
        "concentration_m": chemical.value("concentration_m"),
        "mass_mg": chemical.value("mass_mg"),
        "amount_mmol": chemical.value("amount_mmol"),
        "volume_ml": chemical.value("volume_ml"),
        "quantities": [asdict(quantity) for quantity in chemical.quantities],
        "warnings": chemical.warnings,
        "property_sources": chemical.property_sources,
        "catalog_status": chemical.catalog_status,
        "formulation": chemical.formulation,
    }


def _build_template_text(text: str, chemicals: list[ParsedChemical]) -> str:
    replacements: list[tuple[int, int, str]] = []
    occupied: list[tuple[int, int]] = []
    for chemical in chemicals:
        if chemical.role == "workup":
            continue
        if chemical.variable:
            for start, end in chemical.name_spans:
                replacements.append((start, end, f"{{{chemical.alias}.name}}"))
                occupied.append((start, end))
        for quantity in chemical.quantities:
            if quantity.derived or quantity.span == (-1, -1):
                continue
            replacement = _quantity_alias(chemical.alias, quantity)
            if replacement:
                replacements.append((*quantity.span, replacement))

    output = text
    for start, end, replacement in sorted(replacements, key=lambda item: (item[0], item[1]), reverse=True):
        if any(other_start <= start and end <= other_end for other_start, other_end in occupied if (other_start, other_end) != (start, end)):
            continue
        output = output[:start] + replacement + output[end:]
    return output


def _quantity_alias(alias: str, quantity: ParsedQuantity) -> str | None:
    unit = quantity.source_unit.casefold().rstrip(".")
    if quantity.kind == "mass_mg":
        if unit == "kg":
            return f"{{{alias}.kg}} kg"
        if unit == "g":
            return f"{{{alias}.g}} g"
        return f"{{{alias}.mg}} mg"
    if quantity.kind == "amount_mmol":
        if unit == "mol":
            return f"{{{alias}.mol}} mol"
        return f"{{{alias}.mmol}} mmol"
    if quantity.kind == "volume_ml":
        if unit in {"µl", "μl", "ul"}:
            return f"{{{alias}.mcl}} µL"
        if unit == "l":
            return f"{{{alias}.l}} L"
        return f"{{{alias}.ml}} mL"
    if quantity.kind == "equivalents":
        return f"{{{alias}.eq}} equiv"
    # Concentrations and percentages are method constants; the renderer has no
    # corresponding dynamic alias and should preserve them verbatim.
    return None


def _float_or_none(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _format_number(value: Any) -> str:
    number = _float_or_none(value)
    if number is None:
        return ""
    return f"{number:.6g}"
