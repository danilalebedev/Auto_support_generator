from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu
from PIL import Image

from ..chemdraw_ole import insert_chemdraw_marker_objects


BOOKMARK = "asg_scope_overview"


def insert_overview(docx_path, model_path):
    """Insert editable ChemDraw scope pages before the characterization section."""
    document = Document(str(docx_path))
    body = document._element.body
    remove_overview(body)
    model_path = Path(model_path)
    model = json.loads(model_path.read_text(encoding="utf-8"))
    section = document.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin
    max_height = section.page_height - section.top_margin - section.bottom_margin - 200000
    nodes = []
    ole_markers: dict[str, Path] = {}
    ole_sizes: dict[str, tuple[float, float]] = {}
    for index, page in enumerate(model["pages"], start=1):
        picture = model_path.parent / page["png"]
        with Image.open(picture) as image:
            ratio = image.height / image.width
        width = min(max_width, int(max_height / ratio))
        height = int(width * ratio)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        ole_source = page.get("cdx") or page.get("cdxml")
        if ole_source:
            marker_text = f"[[ASG_SCOPE_OLE:{index}]]"
            paragraph.add_run(marker_text)
            ole_markers[marker_text] = model_path.parent / ole_source
            ole_sizes[marker_text] = (Emu(width).pt, Emu(height).pt)
        else:
            paragraph.add_run().add_picture(str(picture), width=Emu(width))
        nodes.append(paragraph._p)
        page_break = document.add_page_break()
        nodes.append(page_break._p)
    if not nodes:
        document.save(str(docx_path))
        return
    ids = [int(e.get(qn("w:id"))) for e in body.iter(qn("w:bookmarkStart"))]
    marker = OxmlElement("w:bookmarkStart")
    marker.set(qn("w:id"), str(max(ids, default=0) + 1))
    marker.set(qn("w:name"), BOOKMARK)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), marker.get(qn("w:id")))
    nodes[0].insert(1 if nodes[0].find(qn("w:pPr")) is not None else 0, marker)
    nodes[-1].append(end)
    for index, node in enumerate(nodes):
        body.insert(index, node)
    document.save(str(docx_path))
    if ole_markers:
        insert_chemdraw_marker_objects(docx_path, ole_markers, size_points=ole_sizes)


def remove_overview(body):
    active = False
    marker_id = None
    for node in list(body):
        for start in node.iter(qn("w:bookmarkStart")):
            if start.get(qn("w:name")) == BOOKMARK:
                active = True
                marker_id = start.get(qn("w:id"))
        if active:
            finished = any(e.get(qn("w:id")) == marker_id for e in node.iter(qn("w:bookmarkEnd")))
            body.remove(node)
            if finished:
                break
