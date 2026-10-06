from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

from si_generator.domain.loadings_workflow import read_characterization_template, read_reaction_schema
from si_generator.procedure_import import (
    generate_procedure_inputs,
    generate_procedure_inputs_from_docx,
    parse_procedure,
    read_procedure_text,
)
from si_generator.unified_word_input import materialize_unified_input


PROCEDURE = (
    "To a solution of bromide 2a (400 mg, 1.57 mmol, 1.0 equiv) in CHCl3 "
    "(20.9 mL) were added NBS (279 mg, 1.57 mmol, 1.0 equiv) and DBP "
    "(19.0 mg, 0.078 mmol, 5 mol%). The mixture was stirred for 2 h at 60 °C."
)


class ProcedureImportTests(unittest.TestCase):
    def test_parses_prefix_loading_without_duplicate_overlapping_entities(self) -> None:
        parsed = parse_procedure(
            "1.0 equiv (0.50 mmol) of substrate A was mixed with 2.0 equiv (1.00 mmol) of Et3N "
            "in THF (5.0 mL).",
            variable_names=["substrate A"],
        )
        by_alias = {chemical["alias"]: chemical for chemical in parsed["chemicals"] if chemical["alias"]}

        self.assertEqual(parsed["reference_chemical"], "substrate A")
        self.assertEqual(by_alias["Reagent_1"]["amount_mmol"], 0.5)
        self.assertEqual(by_alias["Et3N"]["equivalents"], 2.0)
        self.assertIn("{Reagent_1.eq} equiv ({Reagent_1.mmol} mmol) of {Reagent_1.name}", parsed["template_text"])
        self.assertIn("{Et3N.eq} equiv ({Et3N.mmol} mmol) of Et3N", parsed["template_text"])
        self.assertNotIn("equiv.mmol", parsed["template_text"])

    def test_respects_explicit_variable_order(self) -> None:
        parsed = parse_procedure(
            "substrate A (100 mg, 1.0 mmol) and substrate B (220 mg, 2.0 mmol) were combined.",
            variable_names=["substrate B", "substrate A"],
        )
        aliases = {chemical["name"]: chemical["alias"] for chemical in parsed["chemicals"]}

        self.assertEqual(aliases["substrate B"], "Reagent_1")
        self.assertEqual(aliases["substrate A"], "Reagent_2")

    def test_derives_stock_solution_amount_before_equivalents(self) -> None:
        parsed = parse_procedure(
            "substrate A (100 mg, 1.0 mmol) was treated with n-BuLi (2.0 M, 1.0 mL).",
            variable_names=["substrate A"],
        )
        by_name = {chemical["name"]: chemical for chemical in parsed["chemicals"]}

        self.assertAlmostEqual(by_name["n-BuLi"]["amount_mmol"], 2.0)
        self.assertAlmostEqual(by_name["n-BuLi"]["equivalents"], 2.0)

    def test_scope_has_mass_column_only_for_limiting_reagent(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            generated = generate_procedure_inputs(
                "substrate A (100 mg, 1.0 mmol) and substrate B (220 mg, 2.0 mmol) were combined.",
                tmp,
                variable_names=["substrate A", "substrate B"],
            )
            table = Document(generated.scope_draft).tables[0]
            headers = [cell.text for cell in table.rows[0].cells]

        self.assertEqual(
            headers,
            [
                "Reagent_1",
                "Mass of Reagent_1, mg",
                "Reagent_2",
                "Product",
                "Product_number",
                "Mass of product, mg",
            ],
        )

    def test_keeps_workup_quantities_literal_and_out_of_schema(self) -> None:
        procedure = (
            "substrate A (100 mg, 1.0 mmol) was stirred in THF (10 mL), then quenched with water "
            "(20 mL) and extracted with EtOAc (30 mL)."
        )
        with tempfile.TemporaryDirectory() as tmp:
            parsed = parse_procedure(procedure, variable_names=["substrate A"])
            generated = generate_procedure_inputs(procedure, tmp, variable_names=["substrate A"])
            schema = read_reaction_schema(generated.reaction_schema)

        self.assertEqual(parsed["ignored_workup_entities"], ["water", "EtOAc"])
        self.assertIn("water (20 mL)", parsed["template_text"])
        self.assertIn("EtOAc (30 mL)", parsed["template_text"])
        self.assertNotIn("Solvent_water", schema)
        self.assertNotIn("Solvent_EtOAc", schema)

    def test_parses_equivalent_method_into_existing_alias_contract(self) -> None:
        parsed = parse_procedure(PROCEDURE, variable_names=["bromide 2a"])
        by_alias = {chemical["alias"]: chemical for chemical in parsed["chemicals"]}

        self.assertEqual(parsed["reference_chemical"], "bromide 2a")
        self.assertAlmostEqual(parsed["reference_amount_mmol"], 1.57)
        self.assertTrue(by_alias["Reagent_1"]["variable"])
        self.assertEqual(by_alias["NBS"]["equivalents"], 1.0)
        self.assertAlmostEqual(by_alias["DBP"]["equivalents"], 0.05)
        self.assertAlmostEqual(by_alias["Solvent_CHCl3"]["concentration_m"], 0.07512, places=5)
        self.assertIn("{Reagent_1.name}", parsed["template_text"])
        self.assertIn("{Solvent_CHCl3.ml} mL", parsed["template_text"])
        self.assertIn("5 mol%", parsed["template_text"])

    def test_generates_word_inputs_compatible_with_loadings_reader(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            generated = generate_procedure_inputs(
                PROCEDURE,
                tmp,
                variable_names=["bromide 2a"],
                product_numbers=["2a", "2b"],
            )
            schema = read_reaction_schema(generated.reaction_schema)
            template = read_characterization_template(generated.si_template)
            scope_document = Document(generated.scope_draft)
            compound_document = Document(generated.compound_table)
            loadings_document = Document(generated.loadings_table)
            full_template_text = "\n".join(
                paragraph.text for paragraph in Document(generated.si_template).paragraphs
            )
            headers = [cell.text for cell in scope_document.tables[0].rows[0].cells]
            report = json.loads(generated.report.read_text(encoding="utf-8"))
            bundle = materialize_unified_input(generated.all_in_one, Path(tmp) / "materialized")

        self.assertEqual(schema["Reagent_1"].equivalents, 1.0)
        self.assertAlmostEqual(schema["NBS"].mw or 0, 177.98, places=2)
        self.assertAlmostEqual(schema["Solvent_CHCl3"].concentration_M or 0, 0.0751196, places=5)
        self.assertEqual(
            headers,
            [
                "Reagent_1",
                "Mass of Reagent_1, mg",
                "Product",
                "Product_number",
                "Mass of product, mg",
            ],
        )
        self.assertEqual(scope_document.tables[0].cell(1, 3).text, "2a")
        self.assertEqual(scope_document.tables[0].cell(2, 3).text, "2b")
        self.assertEqual(generated.scope.name, "Scope.docx")
        self.assertEqual(generated.all_in_one.name, "All_in_one_input.docx")
        self.assertEqual(
            [cell.text for cell in compound_document.tables[0].rows[0].cells],
            ["number", "structure", "color", "mp", "Rf", "HRMS", "Elemental_analysis", "IR"],
        )
        self.assertEqual(compound_document.tables[0].cell(1, 0).text, "2a")
        self.assertEqual(compound_document.tables[0].cell(2, 0).text, "2b")
        self.assertTrue(bundle.has_complete_loadings)
        self.assertIsNotNone(bundle.si_template)
        self.assertIsNone(scope_document.styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr")))
        self.assertIn("{NBS.mg}", template)
        self.assertIn("{Product.yield.percent}", template)
        self.assertIn("{nmr.1h.peaks}", full_template_text)
        self.assertNotIn("Generated SI method template", template)
        self.assertTrue(report["requires_user_confirmation"])
        loadings = loadings_document.tables[0]
        self.assertGreater(loadings_document.sections[0].page_width, loadings_document.sections[0].page_height)
        self.assertIsNotNone(loadings.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")))
        header_shading = loadings.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn("w:shd"))
        self.assertEqual(header_shading.get(qn("w:fill")), "1F4E78")
        self.assertEqual(
            [cell.text for cell in loadings.rows[0].cells],
            [
                "Alias", "Compound", "Role / scope", "Source loading", "Mass, mg",
                "Amount, mmol", "Volume, mL", "equiv.", "Concentration, M",
                "MW, g/mol", "Density, g/mL", "Source / review",
            ],
        )
        by_alias = {row.cells[0].text: row.cells for row in loadings.rows[1:]}
        self.assertEqual(by_alias["NBS"][4].text, "279")
        self.assertEqual(by_alias["NBS"][5].text, "1.57")
        self.assertEqual(by_alias["NBS"][7].text, "1")
        self.assertEqual(by_alias["Solvent_CHCl3"][6].text, "20.9")
        self.assertIn("calculated concentration", by_alias["Solvent_CHCl3"][11].text)
        variable_shading = by_alias["Reagent_1"][1]._tc.get_or_add_tcPr().find(qn("w:shd"))
        self.assertIsNotNone(variable_shading)
        self.assertEqual(variable_shading.get(qn("w:fill")), "FFF2CC")

    def test_loadings_table_highlights_required_values_missing_from_method(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            generated = generate_procedure_inputs(
                "substrate A (1.0 mmol) was stirred in THF.",
                tmp,
                variable_names=["substrate A"],
            )
            table = Document(generated.loadings_table).tables[0]
            by_alias = {row.cells[0].text: row.cells for row in table.rows[1:]}

        mass_cell = by_alias["Reagent_1"][4]
        volume_cell = by_alias["Solvent_THF"][6]
        self.assertEqual(mass_cell.text, "")
        self.assertEqual(volume_cell.text, "")
        self.assertEqual(mass_cell._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill")), "FFF2CC")
        self.assertEqual(volume_cell._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill")), "FFF2CC")

    def test_reads_procedure_from_docx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "method.docx"
            document = Document()
            document.add_paragraph(PROCEDURE)
            document.save(path)
            loaded = read_procedure_text(path)

        self.assertEqual(loaded, PROCEDURE)

    def test_docx_reader_ignores_heading_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "method.docx"
            document = Document()
            document.add_heading("General procedure", level=1)
            document.add_paragraph(PROCEDURE)
            document.save(path)

            loaded = read_procedure_text(path)
            parsed = parse_procedure(loaded, variable_names=["bromide 2a"])

        self.assertEqual(loaded, PROCEDURE)
        self.assertFalse(parsed["template_text"].startswith("General procedure"))

    def test_docx_entry_point_copies_source_method_into_reproducible_package(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            method = root / "source_method.docx"
            document = Document()
            document.add_paragraph(PROCEDURE)
            document.save(method)

            generated = generate_procedure_inputs_from_docx(
                method,
                root / "package",
                variable_names=["bromide 2a"],
                product_numbers=["2a"],
            )

            self.assertEqual(generated.source_method, root / "package" / "Method.docx")
            self.assertTrue(generated.source_method.exists())
            self.assertEqual(read_procedure_text(generated.source_method), PROCEDURE)

    def test_docx_entry_point_rejects_other_source_formats(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "method.txt"
            source.write_text(PROCEDURE, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "must be a .docx"):
                generate_procedure_inputs_from_docx(source, Path(tmp) / "package")


if __name__ == "__main__":
    unittest.main()
