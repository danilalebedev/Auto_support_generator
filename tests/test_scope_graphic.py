import copy
from contextlib import nullcontext
import json
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree as ET
from PIL import Image
from rdkit import Chem
from rdkit.Chem import rdDepictor

from si_generator.domain.compound import Compound
from si_generator.domain.loadings_workflow import SchemaEntry, ScopeRow
from si_generator.domain.requests import GenerateSIRequest
from si_generator.graph.nodes.scope_graphic import generate_scope_graphic_node
from si_generator.scope_graphic.document import insert_overview
from si_generator.scope_graphic.layout import Canvas, align, copy_ring_fills, make_page, parse, transform
from si_generator.scope_graphic.lifecycle import update_model, refresh_patch
from si_generator.method_selectors import parse_selector
from si_generator.scope_graphic.generator import generate, write_model
from si_generator.structure_metadata import StructureMetadata
from si_generator.gui import _build_generate_request


def xml(smiles="Cc1ccccc1"):
    mol = Chem.MolFromSmiles(smiles)
    rdDepictor.Compute2DCoords(mol)
    root = parse(Chem.MolToCDXMLBlock(mol))
    # RDKit's writer leaves BondLength empty in this build; its output uses
    # 28.8-point bonds. Native ChemDraw exports include this document property.
    root.set("BondLength", "28.8")
    return ET.tostring(root, encoding="unicode")


def model():
    products = [{"id": f"cmp_{i}", "number": f"2{letter}", "yield": f"{70+i}%",
                 "cdxml": xml(), "reactants": [xml("c1ccccc1")]}
                for i, letter in enumerate("abc")]
    return {"products": products, "title": "Scope", "conditions": "toluene, 80 C",
            "reaction": {"number": "2a", "product": xml(), "reactants": [xml("c1ccccc1")], "reagents": ["NBS"]}}


def test_disabled_scope_does_not_open_chemdraw_or_read_inputs(tmp_path):
    request = GenerateSIRequest(tmp_path / "missing.docx", "word", tmp_path / "out.docx")
    with patch("si_generator.scope_graphic.generator.generate") as generate:
        assert generate_scope_graphic_node({"request": request}) == {}
        generate.assert_not_called()


def test_enabled_scope_missing_files_is_actionable_error(tmp_path):
    request = GenerateSIRequest(tmp_path / "missing.docx", "word", tmp_path / "out.docx", show_scope=True)
    result = generate_scope_graphic_node({"request": request})
    assert result["status"] == "fail"
    assert "Scope.docx" in result["issues"][0]["message"]


@pytest.mark.parametrize("mode", ["separate", "all_in_one"])
def test_gui_passes_scope_settings_without_enabling_loadings(tmp_path, mode):
    table = tmp_path / "Compound_table.docx"
    scope = tmp_path / "Scope.docx"
    schema = tmp_path / "Reaction_schema.docx"
    for path in (table, scope, schema):
        path.touch()
    request = _build_generate_request(
        input_kind="word", input_mode=mode, input_path_text=str(table), unified_input_text=str(table),
        output_docx_text=str(tmp_path / "support.docx"), show_scope=True,
        scope_conditions="  CHCl3, 25 C, 2 h  ", scope_title="  Test scope  ",
        loadings_scope_text=str(scope), loadings_schema_text=str(schema), generate_loadings=False,
    )
    assert request.show_scope and not request.generate_loadings
    assert request.scope_title == "Test scope"
    assert request.scope_conditions == "CHCl3, 25 C, 2 h"
    assert request.loadings_scope_docx == (scope if mode == "separate" else None)


def test_long_scope_is_split_without_losing_products(tmp_path):
    data = model()
    data["products"] = [{**copy.deepcopy(data["products"][0]), "number": str(i), "id": f"cmp_{i}"}
                        for i in range(25)]
    def preview(app, cdxml, output):
        Image.new("RGB", (800, 1000), "white").save(output)
    with patch("si_generator.scope_graphic.generator.render_png", side_effect=preview):
        saved = write_model(data, tmp_path, app=object())
    pages = json.loads(saved.read_text())["pages"]
    assert len(pages) == 3
    labels = []
    for page in pages:
        root = parse((tmp_path / page["cdxml"]).read_text())
        labels += [e.text for e in root.findall("page/t/s") if e.text.isdigit()]
    assert labels == [str(i) for i in range(25)]


def test_multiple_reaction_schemas_use_only_their_reagent_columns(tmp_path):
    compounds = [
        Compound(number="4a", name="First", id="first", yield_text="70%"),
        Compound(number="5", name="Second", id="second", yield_text="80%"),
        Compound(number="4b", name="Third", id="third", yield_text="75%"),
    ]
    rows = [
        ScopeRow(
            product_number="4a",
            product_mass_mg=10,
            product=StructureMetadata(),
            reagent_cells={"Reagent_1": (1, 2, 1), "Reagent_2": (1, 2, 3)},
            product_cell=(1, 2, 5),
        ),
        ScopeRow(
            product_number="5",
            product_mass_mg=10,
            product=StructureMetadata(),
            reagent_cells={"Reagent_1": (1, 3, 1), "Reagent_2": (1, 3, 3)},
            product_cell=(1, 3, 5),
        ),
        ScopeRow(
            product_number="4b",
            product_mass_mg=10,
            product=StructureMetadata(),
            reagent_cells={"Reagent_1": (1, 4, 1), "Reagent_2": (1, 4, 3)},
            product_cell=(1, 4, 5),
        ),
    ]
    schemas = [
        (
            parse_selector("4a-4b"),
            {
                "Reagent_1": SchemaEntry("Reagent_1", "Reagent_1", 1),
                "Reagent_2": SchemaEntry("Reagent_2", "Reagent_2", 1),
            },
        ),
        (
            parse_selector("5"),
            {"Reagent_1": SchemaEntry("Reagent_1", "Reagent_1", 1)},
        ),
    ]
    cdx = {
        (1, 2, 1): "r1a",
        (1, 2, 3): "r2a",
        (1, 2, 5): "p1",
        (1, 3, 1): "r1b",
        (1, 3, 5): "p2",
        (1, 4, 1): "r1c",
        (1, 4, 3): "r2c",
        (1, 4, 5): "p3",
    }
    captured = {}

    def save_model(data, output, app):
        captured.update(data)
        return Path(output) / "scope_graphic.json"

    with (
        patch("si_generator.scope_graphic.generator.read_scope", return_value=rows),
        patch("si_generator.scope_graphic.generator.read_reaction_schemas", return_value=schemas),
        patch("si_generator.scope_graphic.generator._extract_cdx_by_cell", return_value=cdx),
        patch("si_generator.scope_graphic.generator.chemdraw", return_value=nullcontext(object())),
        patch("si_generator.scope_graphic.generator.cdx_to_xml", side_effect=lambda _app, value: value),
        patch("si_generator.scope_graphic.generator.write_model", side_effect=save_model),
    ):
        generate(compounds, tmp_path / "Scope.docx", tmp_path / "Reaction_schema.docx", tmp_path)

    assert [len(series["reaction"]["reactants"]) for series in captured["series"]] == [2, 1]
    assert [len(series["products"][0]["reactants"]) for series in captured["series"]] == [2, 1]
    assert [[product["number"] for product in series["products"]] for series in captured["series"]] == [
        ["4a", "4b"],
        ["5"],
    ]


def test_cdxml_contains_unique_ids_labels_yields_and_unchanged_chemistry():
    data = model()
    output, metrics = make_page(data["products"], data["reaction"], data["conditions"], data["title"])
    root = parse(output)
    identifiers = [e.get("id") for e in root.iter() if e.get("id")]
    # Fonts have their own ID namespace.
    identifiers = [e.get("id") for e in root.find("page").iter() if e.get("id")]
    assert len(identifiers) == len(set(identifiers))
    labels = {"".join(run.text or "" for run in e.findall("s")): e for e in root.findall("page/t")}
    expected = ("2a, 70%", "2b, 71%", "2c, 72%")
    assert len({labels[label].get("p").split()[1] for label in expected}) == 1
    for label in expected:
        runs = labels[label].findall("s")
        assert [run.text for run in runs] == [label[:2], label[2:]]
        assert [run.get("face") for run in runs] == ["1", "0"]
    assert "Representative reaction: 2a" not in labels
    assert metrics["columns"] == 3
    mols = Chem.MolsFromCDXML(output)
    assert sorted(Chem.MolToSmiles(m) for m in mols) == sorted(["Cc1ccccc1"] * 4 + ["c1ccccc1"])


def test_rotation_preserves_stereochemistry():
    canvas = Canvas()
    ref = canvas.drawing(xml("C[C@H](O)c1ccccc1"))
    target = canvas.drawing(xml("C[C@H](O)c1ccccc1"))
    transform(target.group, np.array([[0, -1], [1, 0]]), np.array([140, 50]))
    assert len(align(ref, target)) >= 6
    result = Chem.MolsFromCDXML(canvas.xml(500, 500))
    assert len(set(Chem.MolToSmiles(m, isomericSmiles=True) for m in result)) == 1


def test_ring_fill_references_target_atoms_and_survives_layout():
    canvas = Canvas()
    reference = canvas.drawing(xml("c1ccccc1"))
    target = canvas.drawing(xml("Cc1ccccc1"))
    fill = ET.SubElement(reference.group.find("fragment"), "ColoredMolecularArea",
                         id=str(canvas.next_id()), BasisObjects=" ".join(b.get("id") for b in reference.group.iter("b")), bgcolor="3")
    mapping = align(reference, target)
    assert copy_ring_fills(reference, target, mapping, canvas.next_id) == 1
    copied = target.group.find(".//ColoredMolecularArea")
    assert copied is not None
    assert copied.getparent().tag == "fragment"
    target_ids = {e.get("id") for e in target.group.iter()}
    assert set(copied.get("BasisObjects").split()) <= target_ids
    assert copied.get("BasisObjects") != fill.get("BasisObjects")


def test_legacy_ring_fill_spelling_is_normalized_for_native_chemdraw():
    canvas = Canvas()
    drawing = canvas.drawing(xml("c1ccccc1"))
    ET.SubElement(drawing.group.find("fragment"), "coloredmoleculararea",
                  id=str(canvas.next_id()), bgcolor="3",
                  BasisObjects=" ".join(b.get("id") for b in drawing.group.iter("b")))
    native = Canvas()
    result = native.drawing(canvas.xml(300, 300))
    assert result.group.find(".//ColoredMolecularArea") is not None
    assert result.group.find(".//coloredmoleculararea") is None


def test_scope_update_follows_ids_for_swap_remove_and_reorder():
    original = model()
    manifest = {"order": ["cmp_2", "cmp_0"], "compounds": {
        "cmp_0": {"number": "5b"}, "cmp_2": {"number": "5a"}}}
    updated = update_model(original, manifest)["series"][0]
    assert [(p["id"], p["number"], p["yield"]) for p in updated["products"]] == [
        ("cmp_2", "5a", "72%"), ("cmp_0", "5b", "70%")]
    assert updated["reaction"]["number"] == "5a"
    assert original["products"][0]["number"] == "2a"


def test_insertion_is_idempotent_and_precedes_characterization(tmp_path):
    path = tmp_path / "support.docx"
    document = Document()
    document.add_paragraph("Compound descriptions")
    document.save(path)
    image = tmp_path / "scope.png"
    Image.new("RGB", (800, 450), "white").save(image)
    data = tmp_path / "scope_graphic.json"
    data.write_text(json.dumps({"pages": [{"png": image.name}]}))
    insert_overview(path, data)
    insert_overview(path, data)
    result = Document(path)
    assert len(result.inline_shapes) == 1
    assert result.paragraphs[-1].text == "Compound descriptions"
    markers = [e.get(qn("w:name")) for e in result._element.iter(qn("w:bookmarkStart"))]
    assert markers == ["asg_scope_overview"]


def test_patch_without_overview_does_not_use_chemdraw(tmp_path):
    with patch("si_generator.scope_graphic.lifecycle.write_model") as write:
        refresh_patch({}, {}, tmp_path, tmp_path / "missing.docx", tmp_path)
        write.assert_not_called()
