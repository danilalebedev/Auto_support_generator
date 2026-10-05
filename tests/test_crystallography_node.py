from __future__ import annotations

import json
from pathlib import Path
import shutil
import zipfile

from docx import Document
import pytest

from si_generator.crystallography.service import discover_source
from si_generator.crystallography.fields import field_value
from si_generator.crystallography.model import read_cif
from si_generator.domain.compound import Compound
from si_generator.domain.requests import GenerateSIRequest, PatchSIRequest, AddCompoundsRequest
from si_generator.graph.compound_store import make_compound_store
from si_generator.graph.nodes.crystallography import process_crystallography_node
from si_generator.workflows.generate_si import run_generate_si
from si_generator.workflows.patch_si import run_patch_si
from si_generator.workflows.add_compounds import run_add_compounds


CIF = '''data_test
_chemical_formula_sum 'C2 H6 O'
_chemical_formula_weight 46.07
_cell_length_a 10.0(1)
_cell_length_b 11.0(1)
_cell_length_c 12.0(1)
_cell_angle_alpha 90
_cell_angle_beta 90
_cell_angle_gamma 90
_cell_volume 1320
_cell_formula_units_Z 4
_symmetry_space_group_name_H-M 'P 1'
_diffrn_ambient_temperature 150
_cell_measurement_temperature 293(2)
loop_
_atom_site_label
_atom_site_type_symbol
_atom_site_fract_x
_atom_site_fract_y
_atom_site_fract_z
C1 C 0.1 0.2 0.3
C2 C 0.25 0.2 0.3
O1 O 0.3 0.31 0.3
'''


def write_table(path, number):
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['number', 'name', 'formula']):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, [number, 'ethanol', 'C2H6O']):
        cell.text = value
    doc.save(path)


def cif_source(root, number='1a'):
    folder = root / 'CIF_source' / number
    folder.mkdir(parents=True)
    path = folder / 'sample.cif'
    path.write_text(CIF, encoding='utf-8')
    return folder.parent


def test_subset_and_unrecognised_number(tmp_path):
    source = cif_source(tmp_path)
    assert set(discover_source(source, tmp_path/'stage', {'1a', '1b'})) == {'1a'}
    with pytest.raises(ValueError, match='mismatch'):
        discover_source(source, tmp_path/'stage', {'2a'})


def test_zip_wrapper_and_unsafe_zip(tmp_path):
    archive = tmp_path/'source.zip'
    with zipfile.ZipFile(archive, 'w') as z:
        z.writestr('CIF_source/1a/sample.cif', CIF)
    assert set(discover_source(archive, tmp_path/'stage', {'1a', '1b'})) == {'1a'}
    with zipfile.ZipFile(archive, 'w') as z:
        z.writestr('../1a/sample.cif', CIF)
    with pytest.raises(ValueError, match='Unsafe'):
        discover_source(archive, tmp_path/'unsafe', {'1a'})


def test_node_has_real_rows_image_and_temperature_priority(tmp_path):
    source = cif_source(tmp_path)
    compounds, order = make_compound_store([Compound('1a', 'ethanol'), Compound('1b', 'other')])
    request = GenerateSIRequest(tmp_path/'input.docx', 'word', tmp_path/'output/docx/support.docx', cif_source=source)
    result = process_crystallography_node({'request': request, 'compounds': compounds, 'order': order})
    assert result.get('status') != 'fail', result
    compound = compounds[order[0]]
    data = json.loads(Path(compound.crystallography_data_path).read_text(encoding='utf-8'))
    rows = {row['key']: row['value'] for row in data['records'][0]['rows']}
    assert rows['temperature'] == '150'
    assert rows['a'] == '10.0(1)'
    report = Document(compound.crystallography_report_path)
    assert len(report.inline_shapes) == 1 and len(report.tables) == 1
    assert not compounds[order[1]].crystallography_data_path
    assert any(i['code'] == 'CIF_ELLIPSOID_PLOT_REQUIRED' for i in result['issues'])


def test_bad_cif_fails_instead_of_stub(tmp_path):
    source = cif_source(tmp_path)
    (source/'1a/sample.cif').write_text('data_empty', encoding='utf-8')
    compounds, order = make_compound_store([Compound('1a', 'ethanol')])
    request = GenerateSIRequest(tmp_path/'input.docx', 'word', tmp_path/'output/docx/support.docx', cif_source=source)
    result = process_crystallography_node({'request': request, 'compounds': compounds, 'order': order})
    assert result['status'] == 'fail'
    assert not list((tmp_path/'output').rglob('*.docx'))


def test_multiblock_cif_and_metadata(tmp_path):
    source = cif_source(tmp_path)
    path = source/'1a/sample.cif'
    path.write_text(CIF + '\n' + CIF.replace('data_test', 'data_second'), encoding='utf-8')
    assert len(read_cif(path)) == 2
    assert field_value(read_cif(path)[0], 'a') == '10.0(1)'


def test_workflow_patch_and_reformat_preserve_crystals(tmp_path):
    table = tmp_path/'Compound_table.docx'
    write_table(table, '1a')
    source = cif_source(tmp_path)
    request = GenerateSIRequest(table, 'word', tmp_path/'output/support.docx', cif_source=source,
                                no_extract_nmr=True, no_check_support=True, insert_spectra_as='none')
    result = run_generate_si(request)
    assert Path(result['artifacts']['support_docx']).is_file(), result['issues']
    report = Document(result['artifacts']['support_docx'])
    assert len(report.inline_shapes) == 1 and len(report.tables) == 1
    patched = run_patch_si(PatchSIRequest(Path(result['artifacts']['manifest']), renumber={'1a': '2a'}, output_folder=tmp_path/'patch'))
    assert patched.get('status') != 'fail', patched.get('issues')
    text = '\n'.join(p.text for p in Document(patched['artifacts']['support_docx']).paragraphs)
    assert 'Crystal structure of 2a' in text and 'Crystal structure of 1a' not in text
    for entry in patched['manifest']['compounds'].values():
        assert Path(entry['artifacts']['crystallography_data']).is_file()
    relocated = tmp_path / 'relocated'
    shutil.copytree(Path(patched['artifacts']['support_docx']).parent.parent, relocated)
    shutil.rmtree(Path(patched['artifacts']['support_docx']).parent.parent)
    reformat = run_patch_si(PatchSIRequest(relocated / 'docx' / Path(patched['artifacts']['manifest']).name,
                                          journal_profile_id='acs.joc', output_folder=tmp_path/'reformat'))
    assert reformat.get('status') != 'fail', reformat.get('issues')
    formatted = Document(reformat['artifacts']['support_docx'])
    assert len(formatted.tables) == 1 and len(formatted.inline_shapes) == 1
    assert 'Crystal structure of 2a' in '\n'.join(p.text for p in formatted.paragraphs)
    removed = run_patch_si(PatchSIRequest(Path(reformat['artifacts']['manifest']), remove=('2a',), output_folder=tmp_path/'remove'))
    assert removed.get('status') != 'fail', removed.get('issues')
    assert not Document(removed['artifacts']['support_docx']).tables


def test_add_keeps_distinct_cif_bundles_and_custom_template(tmp_path):
    table = tmp_path / 'Compound_table.docx'
    write_table(table, '1a')
    source = cif_source(tmp_path)
    custom = tmp_path / 'Crystallography_template.docx'
    template = Document()
    template.add_paragraph('Custom X-ray data for {Product.number}')
    template.add_paragraph('{Crystal.description}')
    template.add_paragraph('{Crystal.table}')
    template.save(custom)
    generated = run_generate_si(GenerateSIRequest(table, 'word', tmp_path/'initial/support.docx', cif_source=source,
                                                 crystallography_template_docx=custom, no_extract_nmr=True,
                                                 no_check_support=True, insert_spectra_as='none'))
    write_table(table, '2a')
    new_source = cif_source(tmp_path / 'new', '2a')
    (new_source / '2a/sample.json').write_text(json.dumps({'growth': 'New crystals from ethanol.'}), encoding='utf-8')
    added = run_add_compounds(AddCompoundsRequest(Path(generated['artifacts']['manifest']), table, 'word',
                                                cif_source=new_source, output_folder=tmp_path/'added'))
    assert added['status'] == 'pass', added['issues']
    assert not any(i['code'] == 'MANIFEST_MISSING_ARTIFACT' for i in added['issues'])
    doc = Document(added['artifacts']['support_docx'])
    assert len(doc.tables) == 2 and len(doc.inline_shapes) == 2
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Custom X-ray data for 1a' in text and 'Custom X-ray data for 2a' in text
    assert 'New crystals from ethanol.' in text
    entries = added['manifest']['compounds'].values()
    data = [json.loads(Path(e['artifacts']['crystallography_data']).read_text(encoding='utf-8')) for e in entries]
    assert 'New crystals from ethanol.' not in data[0]['records'][0]['description']
    assert 'New crystals from ethanol.' in data[1]['records'][0]['description']
