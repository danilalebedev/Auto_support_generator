from pathlib import Path
import json

from docx import Document
import pytest

from si_generator.domain.requests import GenerateSIRequest
from si_generator.workflows.generate_si import run_generate_si
from si_generator.workflows.series import series_requests
from tests.test_crystallography_node import write_table, cif_source


def prepare(tmp_path):
    source = tmp_path / 'Multiple_series'
    for folder, number in [('01_A','1a'),('02_B','2a')]:
        directory = source / folder
        directory.mkdir(parents=True)
        write_table(directory / 'Compound_table.docx', number)
        template = Document()
        template.add_paragraph('{Product.name} ({Product.number})')
        template.add_paragraph('Method '+folder)
        template.save(directory/'SI_template.docx')
    cif_source(source/'02_B', '2a')
    return GenerateSIRequest(source, 'word', tmp_path/'out/support.docx', series_folder=source,
                             no_extract_nmr=True, no_check_support=True, insert_spectra_as='none')


def test_different_methods_merge_without_losing_crystal_table(tmp_path):
    request = prepare(tmp_path)
    result = run_generate_si(request)
    assert result['status'] in {'completed', 'completed_with_warnings'}, result.get('issues')
    assert [e['number'] for e in result['manifest']['compounds'].values()] == ['1a','2a']
    doc = Document(result['artifacts']['support_docx'])
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Method 01_A' in text and 'Method 02_B' in text
    assert 'Crystal structure of 2a' in text and len(doc.tables) == 1
    assert len(result['manifest']['series']) == 2
    crystal_issues = [i for i in result['issues'] if i['code'].startswith('CIF_')]
    second_id = next(key for key, entry in result['manifest']['compounds'].items() if entry['number'] == '2a')
    assert crystal_issues and all(i['compound_id'] == second_id for i in crystal_issues)
    summary = json.loads(Path(result['artifacts']['run_summary']).read_text(encoding='utf-8'))
    assert summary['compound_count'] == 2
    for entry in result['manifest']['compounds'].values():
        for relative in entry['relative_artifacts'].values():
            assert (Path(result['artifacts']['output_root'])/relative).exists()


def test_duplicate_numbers_rejected_before_running_spectra(tmp_path):
    request = prepare(tmp_path)
    write_table(request.series_folder/'02_B/Compound_table.docx','1a')
    with pytest.raises(ValueError, match='Duplicate compound number'):
        series_requests(request)
    assert not request.output_path.parent.exists()


def test_incomplete_loadings_pair_is_error(tmp_path):
    request = prepare(tmp_path)
    Document().save(request.series_folder/'01_A/Reaction_schema.docx')
    with pytest.raises(ValueError, match='together'):
        series_requests(request)


def test_each_series_calculates_its_own_loadings(tmp_path):
    request = prepare(tmp_path)
    for folder, number, equivalent in [('01_A', '1a', '2'), ('02_B', '2a', '3')]:
        root = request.series_folder / folder
        schema = Document()
        table = schema.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells, ['Reagents', 'equiv.', 'MW, g/mol']):
            cell.text = value
        for values in [('Reagent_1', '1', '100'), ('NBS', equivalent, '178')]:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        schema.save(root / 'Reaction_schema.docx')
        scope = Document()
        table = scope.add_table(rows=2, cols=4)
        for cell, value in zip(table.rows[0].cells, ['Reagent_1', 'Mass of Reagent_1, mg', 'Product_number', 'Mass of product, mg']):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ['A', '100', number, '42']):
            cell.text = value
        scope.save(root / 'Scope.docx')
        template = Document()
        template.add_paragraph('{Product.name} ({Product.number})')
        template.add_paragraph('NBS ({NBS.mg} mg, {NBS.mmol} mmol).')
        template.save(root / 'SI_template.docx')
    result = run_generate_si(request)
    assert result['status'] != 'fail', result.get('issues')
    text = '\n'.join(p.text for p in Document(result['artifacts']['support_docx']).paragraphs)
    assert 'NBS (356 mg, 2 mmol)' in text
    assert 'NBS (534 mg, 3 mmol)' in text
