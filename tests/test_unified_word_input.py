from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document

from si_generator.domain.requests import GenerateSIRequest
from si_generator.graph.nodes.unified_input import prepare_unified_input_node
from si_generator.unified_word_input import (
    SECTION_MARKERS,
    build_unified_input_docx,
    materialize_unified_input,
)
from si_generator.word_input import read_word_compounds
from si_generator.workflows.generate_si import run_generate_si


REPO_ROOT = Path(__file__).resolve().parents[1]


class UnifiedWordInputTests(unittest.TestCase):
    def test_materializes_all_sections_and_preserves_compound_ole_parts(self) -> None:
        source = REPO_ROOT / "examples" / "example_1" / "All_in_one_input.docx"
        with tempfile.TemporaryDirectory() as tmp:
            bundle = materialize_unified_input(source, Path(tmp) / "parts")
            compounds = read_word_compounds(bundle.compound_table, extract_structure_metadata=False)

            with zipfile.ZipFile(source) as archive:
                source_ole = [name for name in archive.namelist() if "embedding" in name.lower()]
            with zipfile.ZipFile(bundle.compound_table) as archive:
                extracted_ole = [name for name in archive.namelist() if "embedding" in name.lower()]

        self.assertEqual([compound.number for compound in compounds], ["2a", "2b", "2c", "2d"])
        self.assertTrue(bundle.has_complete_loadings)
        self.assertIsNotNone(bundle.si_template)
        self.assertEqual(len(source_ole), len(extracted_ole))
        self.assertGreater(len(extracted_ole), 0)

    def test_optional_sections_may_be_absent(self) -> None:
        compound_table = REPO_ROOT / "examples" / "example_1" / "Compound_table.docx"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            unified = build_unified_input_docx(compound_table, root / "minimal.docx")
            bundle = materialize_unified_input(unified, root / "parts")

        self.assertIsNone(bundle.reaction_schema)
        self.assertIsNone(bundle.scope)
        self.assertIsNone(bundle.si_template)
        self.assertFalse(bundle.has_complete_loadings)

    def test_incomplete_optional_loadings_sections_are_skipped_with_warning(self) -> None:
        example = REPO_ROOT / "examples" / "example_1"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            unified = build_unified_input_docx(
                example / "Compound_table.docx",
                root / "incomplete.docx",
                reaction_schema=example / "Reaction_schema.docx",
            )
            request = GenerateSIRequest(
                input_path=unified,
                input_kind="word",
                output_path=root / "support.docx",
                unified_input_docx=unified,
                generate_loadings=True,
            )
            result = prepare_unified_input_node({"run_id": "test", "request": request, "issues": []})

        resolved = result["request"]
        self.assertFalse(resolved.generate_loadings)
        self.assertIsNone(resolved.loadings_schema_docx)
        self.assertIsNone(resolved.loadings_scope_docx)
        self.assertEqual(result["issues"][0]["code"], "UNIFIED_INPUT_LOADINGS_INCOMPLETE")

    def test_example_contains_human_readable_section_markers(self) -> None:
        document = Document(REPO_ROOT / "examples" / "example_1" / "All_in_one_input.docx")
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for marker in SECTION_MARKERS.values():
            self.assertIn(marker, text)

    def test_minimal_all_in_one_input_runs_without_optional_sections(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            table_path = root / "Compound_table.docx"
            document = Document()
            table = document.add_table(rows=2, cols=4)
            for index, value in enumerate(("number", "name", "color", "mp")):
                table.cell(0, index).text = value
            for index, value in enumerate(("1a", "Example compound", "white solid", "-")):
                table.cell(1, index).text = value
            document.save(table_path)
            unified = build_unified_input_docx(table_path, root / "All_in_one_input.docx")

            result = run_generate_si(
                GenerateSIRequest(
                    input_path=unified,
                    input_kind="word",
                    output_path=root / "out" / "support_information.docx",
                    unified_input_docx=unified,
                    no_extract_nmr=True,
                    insert_spectra_as="none",
                    no_check_support=True,
                )
            )
            self.assertEqual(result["status"], "completed_with_warnings")
            self.assertFalse(result["generation_config"]["generate_loadings"])
            self.assertTrue(Path(result["output_path"]).exists())


if __name__ == "__main__":
    unittest.main()
