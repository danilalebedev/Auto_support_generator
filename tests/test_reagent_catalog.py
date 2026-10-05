import json
import tempfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from si_generator.reagent_catalog import lookup_reagent, catalog_summary
from si_generator.procedure_import import parse_procedure, generate_procedure_inputs


@pytest.mark.parametrize("name,mw",[("NBS",177.98),("NaOtBu",96.10),("LiHMDS",167.33),("TEMPO",156.25),("BH3·SMe2",75.97),("K2CO3",138.205)])
def test_core_identity_and_formula_corrections(name,mw):
    result = lookup_reagent(name)
    assert result["status"] == "matched"
    assert result["record"]["molecular_weight_g_mol"] == pytest.approx(mw,rel=.001)
    assert result["record"]["source_url"].startswith("https://")


def test_ambiguous_and_formulated_reagents_are_not_pure_substances():
    for name in ("DBP","ether","petroleum ether","10% Pd/C","RuCl3·xH2O"):
        assert lookup_reagent(name)["record"] is None
    assert lookup_reagent("LiOH.H2O")["record"]["formula"] != lookup_reagent("LiOH")["record"]["formula"]


def test_catalog_size_and_temperature_provenance():
    assert catalog_summary()["compound_count"] >= 15000
    assert catalog_summary()["unreviewed_density_evidence_count"] >= 2900
    et3n = lookup_reagent("Et3N")["record"]
    assert et3n["density_g_ml"] == pytest.approx(.7275)
    assert et3n["density_temperature_c"] == 20
    solid = lookup_reagent("K2CO3")["record"]
    assert solid["density_evidence"]
    assert solid.get("density_g_ml") is None


def test_numbered_substrate_nested_name_and_quantities():
    result = parse_procedure("To a solution of 21 (100 mg, 0.5 mmol) in THF (2 mL) were added Pd(PPh3)4 (5 mol%) and NBS (1.2 equiv).", variable_names=["21"])
    values = {c["name"]:c for c in result["chemicals"]}
    assert values["21"]["alias"] == "Reagent_1"
    assert values["Pd(PPh3)4"]["equivalents"] == .05
    assert values["NBS"]["molecular_weight_g_mol"] == 177.98
    assert values["NBS"]["property_sources"]["molecular_weight_g_mol"]["source"] == "local_catalog"


def test_reaction_and_workup_same_solvent_remain_separate():
    result = parse_procedure("substrate A (100 mg, 1 mmol) was stirred in THF (5 mL), then quenched with water (2 mL) and extracted with THF (20 mL).",variable_names=["substrate A"])
    assert "THF (20 mL)" in result["template_text"]
    assert result["template_text"].count("{Solvent_THF.ml}") == 1


def test_solution_does_not_use_neat_density():
    result = parse_procedure("substrate A (1 mmol) was treated with Et3N (2 M, 0.5 mL).",variable_names=["substrate A"])
    chemical = next(c for c in result["chemicals"] if c["name"] == "Et3N")
    assert chemical["density_g_ml"] is None
    assert chemical["amount_mmol"] == 1


def test_conflicting_source_values_are_preserved_and_flagged():
    result = parse_procedure("substrate A (1 mmol) and NBS (98 mg, 1 mmol, 2 equiv) were combined.",variable_names=["substrate A"])
    chemical = next(c for c in result["chemicals"] if c["name"] == "NBS")
    assert chemical["mass_mg"] == 98
    assert chemical["molecular_weight_g_mol"] == 177.98
    assert any("MW conflict" in i["issue"] for i in result["unresolved"])
    assert any("Equivalents conflict" in i["issue"] for i in result["unresolved"])


def test_unknown_fields_highlight_without_polluting_numeric_cells(tmp_path):
    files = generate_procedure_inputs("substrate A (1 mmol) was treated with reagent Q (2 equiv) and NBS (1 equiv) in THF (5 mL).",tmp_path,variable_names=["substrate A"])
    schema = Document(files.reaction_schema).tables[0]
    unknown = next(row for row in schema.rows if row.cells[0].text == "reagent_Q")
    assert unknown.cells[2].text == ""
    assert unknown.cells[2]._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill")) == "FFF2CC"
    template = Document(files.si_template)
    highlighted = [run.text for p in template.paragraphs for run in p.runs if run.font.highlight_color]
    assert "{Reagent_1.name}" in highlighted
    assert "{reagent_Q.eq}" in highlighted
    assert "{NBS.eq}" not in highlighted


def test_unknown_literal_catalyst_is_highlighted(tmp_path):
    files = generate_procedure_inputs("substrate A (1 mmol) was treated with catalyst Q (5 mol%).",tmp_path,variable_names=["substrate A"])
    highlighted = [run.text for p in Document(files.si_template).paragraphs for run in p.runs if run.font.highlight_color]
    assert "catalyst Q" in highlighted
