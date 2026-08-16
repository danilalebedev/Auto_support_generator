from __future__ import annotations

import base64
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

from si_generator.docx_builder import build_document_from_model
from si_generator.domain.compound import Compound
from si_generator.domain.journal_validation import validate_compounds_for_journal
from si_generator.domain.requests import GenerateSIRequest
from si_generator.graph.nodes.settings import load_settings_node
from si_generator.journal_profiles import (
    get_journal_profile,
    journal_profile_defaults,
    journal_profile_manifest_block,
    list_journal_profiles,
    journal_profile_labels,
    resolve_journal_profile_id,
    validate_profile_resources,
)
from si_generator.render.document_model import build_si_document_model


class JournalProfileTests(unittest.TestCase):
    def test_acs_joc_and_organic_letters_are_visible_in_gui_labels(self) -> None:
        labels = journal_profile_labels()

        self.assertEqual(labels["ACS - The Journal of Organic Chemistry (JOC)"], "acs.joc")
        self.assertEqual(labels["ACS - Organic Letters (Org. Lett.)"], "acs.orglett")
        self.assertEqual(resolve_journal_profile_id("ACS - Journal of Organic Chemistry"), "acs.joc")
        self.assertEqual(resolve_journal_profile_id("ACS - Organic Letters"), "acs.orglett")

    def test_acs_journal_templates_do_not_embed_a_specific_reaction_method(self) -> None:
        for profile_id in ("acs.joc", "acs.orglett"):
            with self.subTest(profile_id=profile_id):
                document = Document(get_journal_profile(profile_id).template_path)
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)

                self.assertIn("{Product.preparation}", text)
                self.assertNotIn("GP1", text)
                self.assertNotIn("GP2", text)
                self.assertNotIn("{Reagent_", text)

    def test_all_publication_profiles_have_packaged_resources(self) -> None:
        profiles = list_journal_profiles()

        self.assertGreaterEqual(len(profiles), 17)
        self.assertEqual([], [problem for profile in profiles for problem in validate_profile_resources(profile)])

    def test_journal_override_inherits_publisher_defaults(self) -> None:
        profile = get_journal_profile("acs.joc")
        defaults = journal_profile_defaults(profile.id)

        self.assertEqual(profile.publisher, "ACS Publications")
        self.assertEqual(profile.data["document"]["font_name"], "Arial")
        self.assertEqual(defaults["x_range_ppm_1h"], (-1.0, 10.0))
        self.assertTrue(defaults["template_docx"].exists())

    def test_manifest_block_records_sources_and_profile_version(self) -> None:
        block = journal_profile_manifest_block("nature.commschem")

        self.assertEqual(block["id"], "nature.commschem")
        self.assertTrue(block["source_urls"])
        self.assertEqual(block["visual_style_status"], "house_default")

    def test_settings_node_applies_profile_resources_and_records_overrides(self) -> None:
        request = GenerateSIRequest(
            input_path=Path("input.docx"),
            input_kind="word",
            output_path=Path("support.docx"),
            journal_profile_id="acs.joc",
            x_range_ppm_1h=(-2.0, 11.0),
        )

        result = load_settings_node({"request": request})

        self.assertEqual(result["journal_profile"]["id"], "acs.joc")
        self.assertIn("x_range_ppm_1h", result["journal_profile"]["user_overrides"])
        self.assertTrue(request.template_docx and request.template_docx.exists())
        self.assertTrue(request.mnova_graphics_profile_1h and request.mnova_graphics_profile_1h.exists())

    def test_communications_chemistry_requires_fluorine_nmr(self) -> None:
        compound = Compound(
            id="cmp_001",
            number="2a",
            name="Fluorinated example",
            formula="C8H7F",
            h1_nmr="1H NMR",
            c13_nmr="13C NMR",
            hrms_found="123.1000",
        )

        issues = validate_compounds_for_journal([compound], "nature.commschem")

        self.assertIn("JOURNAL_HETERO_NMR_REQUIRED", {issue["code"] for issue in issues})

    def test_communications_chemistry_appendix_uses_landscape_sections(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "2a_1H.png"
            image.write_bytes(base64.b64decode(_ONE_PIXEL_PNG))
            output = root / "support.docx"
            compound = Compound(
                id="cmp_001",
                number="2a",
                name="Example",
                h1_nmr="1H NMR",
                c13_nmr="13C NMR",
                h1_image_path=str(image),
            )
            profile = get_journal_profile("nature.commschem")
            build_document_from_model(
                build_si_document_model([compound], spectra_embed_mode="png"),
                output,
                template_path=profile.template_path,
                render_options=profile.data,
            )
            orientations = _section_orientations(output)

        self.assertIn("landscape", orientations)
        profile = get_journal_profile("nature.commschem")
        document = Document(profile.template_path)
        self.assertAlmostEqual(document.sections[0].page_width.cm, 21.0, places=1)
        self.assertAlmostEqual(document.sections[0].page_height.cm, 29.7, places=1)


def _section_orientations(docx_path: Path) -> list[str]:
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(docx_path, "r") as source:
        root = ET.fromstring(source.read("word/document.xml"))
    result = []
    for section in root.iter(f"{{{namespace}}}sectPr"):
        page_size = section.find(f"{{{namespace}}}pgSz")
        orientation = "portrait"
        if page_size is not None and page_size.attrib.get(f"{{{namespace}}}orient") == "landscape":
            orientation = "landscape"
        result.append(orientation)
    return result


_ONE_PIXEL_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


if __name__ == "__main__":
    unittest.main()
