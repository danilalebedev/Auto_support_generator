from __future__ import annotations

import sys
import re
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from si_generator.journal_profiles import list_journal_profiles  # noqa: E402


SOURCE_TEMPLATE = ROOT / "src" / "si_generator" / "templates" / "SI_template.docx"


def main() -> int:
    if not SOURCE_TEMPLATE.exists():
        raise FileNotFoundError(SOURCE_TEMPLATE)

    for profile in list_journal_profiles(include_hidden=True):
        target = profile.template_path
        target.parent.mkdir(parents=True, exist_ok=True)
        document = Document(SOURCE_TEMPLATE)
        _remove_method_specific_paragraphs(document)
        settings = profile.data.get("document", {})
        _apply_page_settings(document, settings)
        _apply_font_settings(document, settings)
        document.core_properties.subject = f"Auto Support Generator preset: {profile.label}"
        document.core_properties.comments = (
            "Generated house template based on the cited journal requirements. "
            "Journal requirements remain authoritative."
        )
        document.save(target)
        print(target.relative_to(ROOT))
    return 0


def _apply_page_settings(document, settings: dict) -> None:
    margins = settings.get("margins_cm", [2.5, 2.5, 2.5, 2.5])
    if not isinstance(margins, list) or len(margins) != 4:
        margins = [2.5, 2.5, 2.5, 2.5]
    top, right, bottom, left = (float(value) for value in margins)
    for section in document.sections:
        if str(settings.get("page_size") or "").strip().upper() == "A4":
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)
        section.top_margin = Cm(top)
        section.right_margin = Cm(right)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)


def _apply_font_settings(document, settings: dict) -> None:
    font_name = str(settings.get("font_name") or "Times New Roman")
    font_size = Pt(float(settings.get("font_size_pt") or 12))
    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = font_size
    for paragraph in _all_paragraphs(document):
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = font_size


def _remove_method_specific_paragraphs(document) -> None:
    generic_keys = {"product.preparation", "reaction.loadings"}
    for paragraph in list(document.paragraphs):
        if paragraph._p.xpath('.//w:br[@w:type="page"]'):
            break
        keys = {match.group(1).strip().lower() for match in re.finditer(r"\{([^{}]+)\}", paragraph.text)}
        if not keys or keys.issubset(generic_keys):
            continue
        specific = keys - {
            "product.name",
            "product.number",
            "product.structure",
            "product.support.warning",
            "nmr.1h.label",
            "nmr.1h.conditions",
            "nmr.1h.peaks",
            "nmr.13c.label",
            "nmr.13c.conditions",
            "nmr.13c.peaks",
            "nmr.extra",
            "hrms.label",
            "hrms.adduct",
            "hrms.formula",
            "hrms.calculated",
            "hrms.found",
            "anal.label",
            "anal.formula",
            "anal.calculated",
            "anal.found",
            "ir.label",
            "ir.method",
            "ir.peaks",
        }
        if any(_is_loading_key(key) for key in specific):
            paragraph._element.getparent().remove(paragraph._element)


def _is_loading_key(key: str) -> bool:
    if key.startswith(("reagent_", "solvent_")):
        return True
    if key.startswith("product.") and key.split(".", 1)[1] in {
        "mg", "g", "kg", "mmol", "mol", "yield.percent", "appearance", "mp", "rf.value", "rf.system"
    }:
        return True
    return bool(re.match(r"^[a-z0-9_]+\.(?:name|mg|g|kg|mmol|mol|mcl|ml|l|eq)$", key))


def _all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


if __name__ == "__main__":
    raise SystemExit(main())
