"""Reproduce the eight X-ray compounds from the author's supplied SI and raw data."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
from tempfile import TemporaryDirectory
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / 'src'))
from si_generator.unified_word_input import build_unified_input_docx

# Match by source SI compound and CCDC, not by inconsistent author folder labels.
RECORDS = {
    '3c': (98, 104, 270, '2383313', '3e, 2383313/da8828.cif', 'checkcif_da8828 (1).pdf'),
    '3e': (114, 121, 274, '2535885', 'для депонирования/p-Br (3d)/Da10842_pub.cif', 'Da10842.pdf'),
    '3f': (123, 130, 278, '2432441', '3f, 2432441/da9126a (2).cif', 'checkcif_da9626a.pdf'),
    '3g': (132, 138, 282, '2535886', 'для депонирования/3,4-OMe (3g)/Da10890_pub.cif', 'Da10890.pdf'),
    '3h': (140, 148, 286, '2432443', '3h, 2432443/da9335 (1).cif', 'checkcif_da9335.pdf'),
    '3i': (150, 157, 290, '2535884', 'для депонирования/4-CN (3i)/Da9174_pub.cif', 'Da9174.pdf'),
    '3s': (216, 223, 294, '2364657', '3q, 2364657/da8812.cif', 'checkcif_da8812 (1).pdf'),
    '3v': (237, 245, 298, '2117516', '3u, 2117516/Olga2107dep.cif', None),
}


def clean(text):
    return ' '.join(text.split())


def set_normal(document):
    for style in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        document.styles[style].font.name = 'Times New Roman'
        document.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.space_after = Pt(6)
    document.styles['Normal'].paragraph_format.line_spacing = 1.15
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)


def make_si_template(path):
    document = Document(REPO / 'src/si_generator/templates/SI_template.docx')
    for p in list(document.paragraphs):
        if '{Reagent_' in p.text or p.text == '{reaction.loadings}':
            p._p.getparent().remove(p._p)
    set_normal(document)
    document.save(path)


def make_crystal_template(path):
    document = Document()
    set_normal(document)
    p = document.add_paragraph('{Product.name} ({Product.number})')
    p.runs[0].bold = True
    p.paragraph_format.keep_with_next = True
    for text in ('{Crystal.description}', '{Crystal.table}', '{Crystal.geometry}'):
        document.add_paragraph(text)
    document.save(path)


def compound_table(source_path, output):
    document = Document(source_path)
    paragraphs = list(document.paragraphs)
    document._body.clear_content()
    document.part._styles_part._element = deepcopy(Document().styles.element)
    set_normal(document)
    section = document.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(16.54), Inches(11.69)
    section.left_margin = section.right_margin = Inches(.6)
    section.top_margin = section.bottom_margin = Inches(.6)
    for p in section.header.paragraphs + section.footer.paragraphs:
        p.clear()
    document.add_heading('Eight compounds with single crystal X-ray data', 0)
    document.add_paragraph('Preparation contains the experimental quantities reported in the source SI, including yields, melting points and Rf. '
                           'These are measured source values, not a recalculated common reaction scheme. '
                           'The main NMR data are generated from Spectra_source.zip. Original extra NMR descriptions are retained below.')
    headers = ['Number', 'Structure', 'Name', 'Preparation', 'HRMS (ESI-TOF) m/z', 'HRMS adduct', 'IR', 'Anal', 'Extra NMR']
    widths = [Inches(v) for v in (.5, 1.3, 2.0, 4.3, .75, .8, 2.0, 1.0, 2.65)]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for number, (start, image, *_rest) in RECORDS.items():
        if not paragraphs[start].text.rstrip().endswith(f'({number})'):
            raise ValueError(f'Source changed: paragraph {start} is not compound {number}')
        name = re.sub(r'\s*\(' + number + r'\)\s*$', '', clean(paragraphs[start].text))
        prep = clean(paragraphs[start + 1].text)
        if number == '3e':
            prep += '\n\n' + clean(paragraphs[116].text)
        method_index = 80 if number == '3v' else (78 if number == '3s' else 77)
        prep += '\n\n' + clean(paragraphs[method_index].text)
        analytical = [clean(p.text) for p in paragraphs[start + 2:image]]
        hrms = next(t for t in analytical if t.startswith('HRMS'))
        found = re.search(r'Found\s*:?\s*([\d.]+)', hrms).group(1)
        adduct = re.search(r'\[M[^]]+\]\+', hrms).group(0)
        ir = next(t for t in analytical if t.startswith('IR'))
        anal = next((t for t in analytical if t.startswith('Anal.')), '-')
        if 'Found:' in anal:
            anal = anal.split('Found:', 1)[1].strip().translate(str.maketrans({'С': 'C', 'Н': 'H'}))
        extra = clean(paragraphs[144].text) if number == '3h' else ''
        if number == '3v':
            extra = '\n'.join(clean(paragraphs[i].text) for i in (241, 242))
        row = table.add_row()
        # A long analytical entry may continue on the next page; it must never be clipped.
        for cell, text in zip(row.cells, [number, '', name, prep, found, adduct, ir, anal, extra]):
            cell.text = text
        for element in paragraphs[start + 1]._p.iter(qn('w:object')):
            cloned = deepcopy(element)
            for shape in cloned.iter('{urn:schemas-microsoft-com:vml}shape'):
                style = shape.get('style', '')
                dims = dict(re.findall(r'(width|height):([\d.]+)pt', style))
                if 'width' in dims and float(dims['width']) > 80:
                    scale = 80 / float(dims['width'])
                    for key, value in dims.items():
                        style = re.sub(key + r':[\d.]+pt', f'{key}:{float(value)*scale:.2f}pt', style)
                    shape.set('style', style)
            row.cells[1].paragraphs[0].add_run()._r.append(cloned)
    for row_index, row in enumerate(table.rows):
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = False
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10)
                    run.font.hidden = False
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row_index == 0
    document.add_heading('Source general procedures', 1)
    for index in (77, 78, 80):
        document.add_paragraph(clean(paragraphs[index].text))
    # The table section is landscape; the embedded SI template keeps portrait output pages.
    final = document.add_section(WD_SECTION.NEW_PAGE)
    final.orientation = WD_ORIENT.PORTRAIT
    final.page_width, final.page_height = Inches(8.27), Inches(11.69)
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source-docx', required=True, type=Path)
    parser.add_argument('--fid-zip', required=True, type=Path)
    parser.add_argument('--cif-root', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    source = Document(args.source_docx)
    make_crystal_template(root / 'Crystallography_template.docx')
    with TemporaryDirectory(prefix='asg_cif_example_') as temporary:
        staging = Path(temporary)
        compound_table(args.source_docx, staging / 'Compound_table.docx')
        make_si_template(staging / 'SI_template.docx')
        build_unified_input_docx(
            staging / 'Compound_table.docx',
            root / 'All_in_one_input.docx',
            si_template=staging / 'SI_template.docx',
            crystallography_template=root / 'Crystallography_template.docx',
        )
    provenance = {'source_document': args.source_docx.name, 'source_sha256': hashlib.sha256(args.source_docx.read_bytes()).hexdigest(), 'compounds': {}}
    for number, (start, image, growth, ccdc, filename, pdf_name) in RECORDS.items():
        cif = args.cif_root / filename
        folder = root / 'CIF_source' / number
        folder.mkdir(parents=True, exist_ok=True)
        shutil.copy2(cif, folder / f'{number}.cif')
        ids = source.paragraphs[image]._p.xpath('.//a:blip/@r:embed')
        if len(ids) != 1:
            raise ValueError(f'Expected one original ORTEP image for {number}')
        (folder / 'ORTEP.png').write_bytes(source.part.related_parts[ids[0]].blob)
        caption = 'ORTEP drawing with thermal ellipsoids at the 50% probability level.'
        if number == '3v':
            caption = 'Minor diastereomer B. ' + caption
        metadata = {'ccdc': ccdc, 'growth': clean(source.paragraphs[growth].text), 'image': 'ORTEP.png',
                    'caption': caption, 'include_geometry': False}
        (folder / f'{number}.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf-8')
        if pdf_name:
            shutil.copy2(cif.parent / pdf_name, folder / f'{number}.checkcif.pdf')
        provenance['compounds'][number] = {'ccdc': ccdc, 'cif_source': filename, 'cif_sha256': hashlib.sha256(cif.read_bytes()).hexdigest(),
                                          'heading_paragraph': start, 'ortep_paragraph': image, 'growth_paragraph': growth}
    with zipfile.ZipFile(args.fid_zip) as source_zip, zipfile.ZipFile(root / 'Spectra_source.zip', 'w', zipfile.ZIP_DEFLATED) as main_zip, zipfile.ZipFile(root / 'Additional_spectra.zip', 'w', zipfile.ZIP_DEFLATED) as extra_zip:
        for item in source_zip.infolist():
            parts = item.filename.split('/')
            if item.is_dir() or len(parts) < 4 or parts[0] != 'fid' or parts[1] not in RECORDS:
                continue
            additional = parts[2] == '19F' or 'minor diastereomer' in parts[2]
            (extra_zip if additional else main_zip).writestr('/'.join(parts[1:]), source_zip.read(item))
    (root / 'provenance.json').write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding='utf-8')
    print(root)


if __name__ == '__main__':
    main()
