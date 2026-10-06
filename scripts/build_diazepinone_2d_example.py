"""Build example 5 from the author-provided benzodiazepinone SI and spectra."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from si_generator.structure_metadata import extract_structure_metadata_by_cell, extract_structure_metadata_by_row
from si_generator.unified_word_input import build_unified_input_docx


COMPOUNDS = (
    {
        "number": "4a",
        "name": "7a,12-Dihydro-5H-benzo[2,3][1,4]diazepino[7,1-a]isoindol-6(7H)-one",
        "product_paragraph": 202,
        "precursor_paragraph": 26,
        "precursor_mass_mg": 100,
        "product_mass_mg": 69,
        "appearance": "white solid",
        "mp": "222-223",
        "rf": "0.56 (petroleum ether : ethyl acetate = 1 : 1)",
        "hrms": "[M+H]+ 251.1183",
    },
    {
        "number": "4b",
        "name": "10-Fluoro-7a,12-dihydro-5H-benzo[2,3][1,4]diazepino[7,1-a]isoindol-6(7H)-one",
        "product_paragraph": 209,
        "precursor_paragraph": 31,
        "precursor_mass_mg": 100,
        "product_mass_mg": 62,
        "appearance": "off-white solid",
        "mp": "239-240",
        "rf": "0.40 (petroleum ether : ethyl acetate = 1 : 1)",
        "hrms": "[M+H]+ 269.1088",
    },
    {
        "number": "4c",
        "name": "9-Fluoro-7a,12-dihydro-5H-benzo[2,3][1,4]diazepino[7,1-a]isoindol-6(7H)-one",
        "product_paragraph": 217,
        "precursor_paragraph": 39,
        "precursor_mass_mg": 100,
        "product_mass_mg": 59,
        "appearance": "off-white solid",
        "mp": "222-223",
        "rf": "0.56 (petroleum ether : ethyl acetate = 1 : 1)",
        "hrms": "[M+H]+ 269.1088",
    },
    {
        "number": "4d",
        "name": "9-Nitro-7a,12-dihydro-5H-benzo[2,3][1,4]diazepino[7,1-a]isoindol-6(7H)-one",
        "product_paragraph": 226,
        "precursor_paragraph": 47,
        "precursor_mass_mg": 100,
        "product_mass_mg": 54,
        "appearance": "off-white solid",
        "mp": "262-263",
        "rf": "0.39 (petroleum ether : ethyl acetate = 1 : 1)",
        "hrms": "[M+H]+ 296.1030",
    },
    {
        "number": "4e",
        "name": "10-Bromo-7a,12-dihydro-5H-benzo[2,3][1,4]diazepino[7,1-a]isoindol-6(7H)-one",
        "product_paragraph": 233,
        "precursor_paragraph": 53,
        "precursor_mass_mg": 100,
        "product_mass_mg": 66,
        "appearance": "off-white solid",
        "mp": "233-234",
        "rf": "0.39 (petroleum ether : ethyl acetate = 1 : 1)",
        "hrms": "[M+H]+ 329.0287",
    },
)

SPECTRA = {
    "4a": {
        "1H": Path("4a/da11215_1H"),
        "13C": Path("4a/da11215_13C"),
        "HSQC": Path("4a/da11215_HSQC"),
        "HMBC": Path("4a/da11215_HMBC"),
    },
    "4b": {
        "1H": Path("4b/da9384_1H"),
        "13C": Path("4b/da9384_13C"),
        "HSQC": Path("4b/da9384_HSQC"),
        "HMBC": Path("4b/da9384_HMBC"),
    },
    "4c": {"1H": Path("4c/da9441/1"), "13C": Path("4c/da9441/2")},
    "4d": {"1H": Path("4d/da9871_1H"), "13C": Path("4d/da9871_13C")},
    "4e": {"1H": Path("4e/da9443/1"), "13C": Path("4e/da9443/2")},
}


def _set_normal(document: Document) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles["Normal"]
    normal.font.size = Pt(14)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def _set_cell_margins(cell, value: int = 45) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _normalize_table_ole_styles(document: Document) -> None:
    for shape in document._element.xpath(
        './/*[local-name()="tbl"]//*[local-name()="object"]//*[local-name()="shape"]'
    ):
        style = str(shape.get("style") or "")
        width = re.search(r"(?:^|;)width:([^;]+)", style)
        height = re.search(r"(?:^|;)height:([^;]+)", style)
        if width and height:
            shape.set("style", f"width:{width.group(1)};height:{height.group(1)}")


def _ole_object(paragraph):
    objects = paragraph._p.xpath(".//w:object")
    if len(objects) != 1:
        raise ValueError(f"Expected one ChemDraw object, found {len(objects)} in: {paragraph.text}")
    return deepcopy(objects[0])


def _prepare_source_document(source_path: Path) -> tuple[Document, list]:
    document = Document(source_path)
    paragraphs = list(document.paragraphs)
    for compound in COMPOUNDS:
        _ole_object(paragraphs[int(compound["product_paragraph"])])
        _ole_object(paragraphs[int(compound["precursor_paragraph"])])
    return document, paragraphs


def make_compound_table(source_path: Path, output: Path) -> list[dict[str, object]]:
    document, paragraphs = _prepare_source_document(source_path)
    document._body.clear_content()
    document.part._styles_part._element = deepcopy(Document().styles.element)
    _set_normal(document)
    section = document.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.45)
    section.top_margin = section.bottom_margin = Inches(0.45)

    headers = ("number", "structure", "color", "mp", "Rf", "HRMS", "Elemental_analysis")
    widths = [Inches(value) for value in (0.55, 2.2, 1.1, 0.8, 2.5, 1.25, 2.2)]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    rows = [dict(compound) for compound in COMPOUNDS]
    for compound in rows:
        row = table.add_row()
        values = (
            compound["number"],
            "",
            compound["appearance"],
            compound["mp"],
            compound["rf"],
            compound["hrms"],
            "-",
        )
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
        row.cells[1].paragraphs[0].add_run()._r.append(
            _ole_object(paragraphs[int(compound["product_paragraph"])])
        )

    for row_index, row in enumerate(table.rows):
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for column_index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if row_index == 0 or column_index in {0, 2, 3, 5}:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.font.hidden = False
                    run.bold = row_index == 0

    final = document.add_section(WD_SECTION.NEW_PAGE)
    final.orientation = WD_ORIENT.PORTRAIT
    final.page_width, final.page_height = Inches(8.27), Inches(11.69)
    final.left_margin = final.right_margin = Inches(0.7)
    final.top_margin = final.bottom_margin = Inches(0.7)
    _normalize_table_ole_styles(document)
    document.save(output)

    metadata = extract_structure_metadata_by_row(output)
    for row_index, compound in enumerate(rows, start=2):
        structure = metadata.get(row_index)
        if structure is None or not structure.formula:
            raise ValueError(f"Could not extract product formula for {compound['number']}")
        compound["formula"] = structure.formula
        compound["smiles"] = structure.smiles
    return rows


def make_reaction_schema(path: Path) -> None:
    document = Document()
    _set_normal(document)
    marker = document.add_paragraph("[AUTO SI: REACTION 4a-4e]")
    marker.runs[0].bold = True
    headers = ("Reagents", "equiv.", "MW, g/mol", "Density, g/ml", "Concentration, M")
    values = (
        ("Reagent_1", "1", "", "", ""),
        ("Phenylenediamine", "3", "108.14", "", ""),
        ("AcOH", "5", "60.05", "1.049", ""),
    )
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for item in values:
        row = table.add_row()
        for cell, text in zip(row.cells, item):
            cell.text = text
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
    document.save(path)


def make_scope(source_path: Path, compounds: list[dict[str, object]], output: Path) -> None:
    document, paragraphs = _prepare_source_document(source_path)
    document._body.clear_content()
    document.part._styles_part._element = deepcopy(Document().styles.element)
    _set_normal(document)
    section = document.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.5)
    section.top_margin = section.bottom_margin = Inches(0.5)

    headers = (
        "Reagent_1",
        "Mass of Reagent_1, mg",
        "Product",
        "Product_number",
        "Mass of product, mg",
    )
    widths = [Inches(value) for value in (2.4, 1.5, 2.6, 1.25, 1.5)]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    for compound in compounds:
        row = table.add_row()
        values = (
            "",
            f"{float(compound['precursor_mass_mg']):g}",
            "",
            compound["number"],
            f"{float(compound['product_mass_mg']):g}",
        )
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
        row.cells[0].paragraphs[0].add_run()._r.append(
            _ole_object(paragraphs[int(compound["precursor_paragraph"])])
        )
        row.cells[2].paragraphs[0].add_run()._r.append(
            _ole_object(paragraphs[int(compound["product_paragraph"])])
        )

    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
    _normalize_table_ole_styles(document)
    document.save(output)

    metadata = extract_structure_metadata_by_cell(output)
    for row_index, compound in enumerate(compounds, start=2):
        reagent = metadata.get((1, row_index, 1))
        product = metadata.get((1, row_index, 3))
        if reagent is None or not reagent.formula:
            raise ValueError(f"Scope is missing Reagent_1 for {compound['number']}")
        if product is None or product.formula != compound["formula"]:
            raise ValueError(f"Scope product differs from compound table for {compound['number']}")


def make_si_template(path: Path) -> None:
    document = Document(REPO / "src/si_generator/templates/SI_template.docx")
    for paragraph in list(document.paragraphs):
        if "{Reagent_" in paragraph.text or paragraph.text == "{reaction.loadings}":
            paragraph._p.getparent().remove(paragraph._p)
    first = document.paragraphs[0]._p
    definitions = (
        ("[AUTO SI: METHOD 4a-4e]", True),
        (
            "Benzodiazepinone {Product.number} was obtained from {Reagent_1.name} "
            "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol), o-phenylenediamine "
            "({Phenylenediamine.mg} mg, {Phenylenediamine.mmol} mmol) and AcOH "
            "({AcOH.mcl} µL) according to GP3. Yield {Product.mg} mg "
            "({Product.yield.percent}); {Product.appearance}; mp {Product.mp} °C. "
            "Rf = {Product.rf.value} ({Product.rf.system}).",
            False,
        ),
        ("[AUTO SI: COMPOUND TEMPLATE]", True),
    )
    for text, bold in definitions:
        paragraph = document.add_paragraph(text)
        paragraph.runs[0].bold = bold
        first.addprevious(paragraph._p)
    _set_normal(document)
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
    structure = next(paragraph for paragraph in document.paragraphs if "{Product.structure}" in paragraph.text)
    structure.paragraph_format.line_spacing = Pt(1)
    for run in structure.runs:
        run.font.size = Pt(1)
    document.save(path)


def make_crystallography_template(path: Path) -> None:
    document = Document()
    _set_normal(document)
    heading = document.add_paragraph("{Product.name} ({Product.number})")
    heading.runs[0].bold = True
    heading.paragraph_format.keep_with_next = True
    document.add_paragraph("{Crystal.description}")
    document.add_paragraph("{Crystal.table}")
    document.add_paragraph("{Crystal.geometry}")
    document.save(path)


def make_spectra_archive(source_root: Path, output: Path) -> dict[str, list[str]]:
    coverage: dict[str, list[str]] = {}
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for number, experiments in SPECTRA.items():
            coverage[number] = []
            for nucleus, relative_source in experiments.items():
                source = source_root / relative_source
                if not source.is_dir():
                    raise FileNotFoundError(f"Missing source experiment: {source}")
                expected = "ser" if nucleus in {"HSQC", "HMBC"} else "fid"
                if not (source / expected).is_file():
                    raise FileNotFoundError(f"{number} {nucleus} has no {expected}: {source}")
                for file_path in sorted(path for path in source.rglob("*") if path.is_file()):
                    target = Path(number) / nucleus / file_path.relative_to(source)
                    archive.write(file_path, target.as_posix())
                coverage[number].append(nucleus)
    return coverage


def _strip_comments(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        entries = {item.filename: source.read(item.filename) for item in source.infolist()}
    for name in list(entries):
        if name.startswith("word/comments"):
            entries.pop(name)
    for name in ("word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"):
        payload = entries.get(name)
        if payload is None:
            continue
        text = payload.decode("utf-8")
        text = re.sub(r"<w:commentRange(?:Start|End)[^>]*/>", "", text)
        text = re.sub(r"<w:r[^>]*>\s*<w:commentReference[^>]*/>\s*</w:r>", "", text)
        text = re.sub(r"<Relationship[^>]+Target=\"comments[^\"]*\"[^>]*/>", "", text)
        text = re.sub(r"<Override[^>]+PartName=\"/word/comments[^\"]*\"[^>]*/>", "", text)
        entries[name] = text.encode("utf-8")
    temporary = path.with_suffix(".comments-cleaned.docx")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as output:
        for name, payload in entries.items():
            output.writestr(name, payload)
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-docx", required=True, type=Path)
    parser.add_argument("--spectra-root", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    compounds = make_compound_table(args.source_docx, root / "Compound_table.docx")
    make_reaction_schema(root / "Reaction_schema.docx")
    make_scope(args.source_docx, compounds, root / "Scope.docx")
    make_si_template(root / "SI_template.docx")
    make_crystallography_template(root / "Crystallography_template.docx")
    build_unified_input_docx(
        root / "Compound_table.docx",
        root / "All_in_one_input.docx",
        reaction_schema=root / "Reaction_schema.docx",
        scope=root / "Scope.docx",
        si_template=root / "SI_template.docx",
        crystallography_template=root / "Crystallography_template.docx",
    )
    for filename in (
        "Compound_table.docx",
        "Reaction_schema.docx",
        "Scope.docx",
        "SI_template.docx",
        "Crystallography_template.docx",
        "All_in_one_input.docx",
    ):
        _strip_comments(root / filename)

    coverage = make_spectra_archive(args.spectra_root, root / "Spectra_source.zip")
    provenance = {
        "source_document": str(args.source_docx.resolve()),
        "source_sha256": hashlib.sha256(args.source_docx.read_bytes()).hexdigest(),
        "spectra_root": str(args.spectra_root.resolve()),
        "compound_order": [compound["number"] for compound in compounds],
        "formulae": {compound["number"]: compound["formula"] for compound in compounds},
        "spectra": coverage,
    }
    (root / "provenance.json").write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(root)


if __name__ == "__main__":
    main()
