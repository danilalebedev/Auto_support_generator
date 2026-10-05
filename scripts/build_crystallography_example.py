"""Rebuild the author-provided CIF example without inventing analytical data."""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
import re
import shutil
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HEADINGS = {'3a':81, '3c':98, '3f':123, '3s':216}
CRYSTALS = {'3c':('3e, 2383313/da8828.cif', 104, 270, '2383313'),
            '3f':('3f, 2432441/da9126a (2).cif', 130, 278, '2432441'),
            '3s':('3q, 2364657/da8812.cif', 223, 294, '2364657')}


def create_template(path):
    d = Document()
    d.add_paragraph('Crystallographic data for {Product.number}', 'Heading 1')
    d.add_paragraph('{Crystal.description}')
    d.add_paragraph('{Crystal.table}')
    d.add_paragraph('{Crystal.geometry}')
    for style in ('Normal','Heading 1'):
        d.styles[style].font.name = 'Times New Roman'
        d.styles[style].font.color.rgb = RGBColor(0,0,0)
        d.styles[style].font.size = Pt(11 if style=='Normal' else 12)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)


def compound_table(source, numbers, path):
    d = Document(source)
    paragraphs = list(d.paragraphs)
    d._body.clear_content()
    defaults = Document()
    d.part._styles_part._element = deepcopy(defaults.styles.element)
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = d.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.left_margin = section.right_margin = Inches(.5)
    section.top_margin = section.bottom_margin = Inches(.5)
    title = d.add_heading('Compound table for the crystallography example', 0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt(20)
    d.add_paragraph('Original preparation and analytical data from the supplied Supporting Information. '
                    '3a deliberately has no CIF. Raw NMR spectra are supplied separately.')
    headers = ['Number','Structure','Name','Preparation','HRMS [M+H]+','Anal']
    table = d.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    widths = [Inches(v) for v in (.7, 1.65, 2.2, 4.05, .9, 1.2)]
    for column, width in zip(table.columns, widths):
        column.width = width
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for number in numbers:
        start = HEADINGS[number]
        row = table.add_row()
        name = re.sub(r'\s*\('+re.escape(number)+r'\)\s*$', '', paragraphs[start].text)
        prep = paragraphs[start+1]
        analytical = []
        for p in paragraphs[start+2:]:
            if re.search(r'\([2356][a-z]?\)\s*$', p.text) and not p.text.startswith(('HRMS','IR','1H','13C','Anal')):
                break
            analytical.append(p.text)
        found = next(re.search(r'Found\s*:?\s*([\d.]+)', t).group(1) for t in analytical if t.startswith('HRMS'))
        anal = next((t for t in analytical if t.startswith('Anal.')), '-')
        if 'Found:' in anal:
            anal = anal.split('Found:', 1)[1].strip().translate(str.maketrans({'С':'C','Н':'H'}))
        for cell, text in zip(row.cells, [number, '', name, prep.text, found, anal]):
            cell.text = text
        for element in prep._p.iter(qn('w:object')):
            row.cells[1].paragraphs[0].add_run()._r.append(deepcopy(element))
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = False
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.font.hidden = False
                    run.bold = row_index == 0
    d.save(path)


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--source-docx', type=Path, required=True)
    p.add_argument('--fid-zip', type=Path, required=True)
    p.add_argument('--cif-root', type=Path, required=True)
    p.add_argument('--output', type=Path, required=True)
    args = p.parse_args()
    root = args.output
    root.mkdir(parents=True, exist_ok=True)
    repo = Path(__file__).resolve().parents[1]
    source = Document(args.source_docx)
    compound_table(args.source_docx, list(HEADINGS), root/'Compound_table.docx')
    shutil.copy2(repo/'src/si_generator/templates/SI_template.docx', root/'SI_template.docx')
    create_template(root/'Crystallography_template.docx')
    create_template(repo/'src/si_generator/templates/Crystallography_template.docx')
    for number, (filename, image_index, growth_index, ccdc) in CRYSTALS.items():
        folder = root/'CIF_source'/number
        folder.mkdir(parents=True, exist_ok=True)
        target = folder/f'{number}.cif'
        shutil.copy2(args.cif_root/filename, target)
        ids = source.paragraphs[image_index]._p.xpath('.//a:blip/@r:embed')
        if len(ids) != 1:
            raise ValueError(f'Expected one original ORTEP image for {number}')
        (folder/'ORTEP.png').write_bytes(source.part.related_parts[ids[0]].blob)
        metadata = {'ccdc':ccdc, 'growth':source.paragraphs[growth_index].text,
                    'image':'ORTEP.png', 'caption':'ORTEP drawing with thermal ellipsoids at the 50% probability level.',
                    'include_geometry':False}
        target.with_suffix('.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf-8')
    # Write a small ZIP directly; keep the supplied archive untouched.
    with zipfile.ZipFile(args.fid_zip) as z, zipfile.ZipFile(root/'Spectra_source.zip','w',zipfile.ZIP_DEFLATED) as out:
        for item in z.infolist():
            parts = Path(item.filename).parts
            if len(parts)>2 and parts[0]=='fid' and parts[1] in HEADINGS and not item.is_dir():
                out.writestr('/'.join(parts[1:]), z.read(item))
    for folder_name, numbers in [('01_Method_A',['3a','3c','3f']),('02_Method_B',['3s'])]:
        folder = root/'Multiple_series'/folder_name
        folder.mkdir(parents=True, exist_ok=True)
        compound_table(args.source_docx, numbers, folder/'Compound_table.docx')
        shutil.copy2(root/'SI_template.docx',folder/'SI_template.docx')
        with zipfile.ZipFile(root/'Spectra_source.zip') as z, zipfile.ZipFile(folder/'Spectra_source.zip','w',zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                if item.filename.split('/')[0] in numbers:
                    out.writestr(item.filename,z.read(item))
        for number in numbers:
            if number in CRYSTALS:
                shutil.copytree(root/'CIF_source'/number,folder/'CIF_source'/number,dirs_exist_ok=True)
    print(root)


if __name__ == '__main__':
    main()
