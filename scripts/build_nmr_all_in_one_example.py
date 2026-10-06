"""Build the complete NMR and crystallography all-in-one example from author data."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
from tempfile import TemporaryDirectory
import time
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from si_generator.structure_metadata import extract_structure_metadata_by_cell, extract_structure_metadata_by_row
from si_generator.unified_word_input import build_unified_input_docx


COMPOUND_NUMBERS = (
    "2g",
    "2h",
    "2k",
    *(f"3{letter}" for letter in "abcdefghijklmnopqrstuvw"),
    "5",
    "6",
)

# Match source-SI identities and CCDC numbers, not the inconsistent source folder labels.
CRYSTAL_RECORDS = {
    "3c": (104, 270, "2383313", "3e, 2383313/da8828.cif", "checkcif_da8828 (1).pdf"),
    "3e": (121, 274, "2535885", "для депонирования/p-Br (3d)/Da10842_pub.cif", "Da10842.pdf"),
    "3f": (130, 278, "2432441", "3f, 2432441/da9126a (2).cif", "checkcif_da9626a.pdf"),
    "3g": (138, 282, "2535886", "для депонирования/3,4-OMe (3g)/Da10890_pub.cif", "Da10890.pdf"),
    "3h": (148, 286, "2432443", "3h, 2432443/da9335 (1).cif", "checkcif_da9335.pdf"),
    "3i": (157, 290, "2535884", "для депонирования/4-CN (3i)/Da9174_pub.cif", "Da9174.pdf"),
    "3s": (223, 294, "2364657", "3q, 2364657/da8812.cif", "checkcif_da8812 (1).pdf"),
    "3v": (245, 298, "2117516", "3u, 2117516/Olga2107dep.cif", None),
}

# The preparation paragraphs for 5 and 6 contain complete reaction schemes. Use the
# dedicated product objects from the spectral appendix so metadata extraction cannot
# select a starting material from the scheme.
SPECTRAL_STRUCTURE_PARAGRAPHS = {
    "5": 579,
    "6": 586,
}


PRECURSOR_SMILES = {
    "1a": "O=C1C(=Cc2ccccc2)C(=O)c2ccccc21",
    "1b": "Cc1ccc(C=C2C(=O)c3ccccc3C2=O)cc1",
    "1c": "Cc1ccccc1C=C1C(=O)c2ccccc2C1=O",
    "1d": "O=C1C(=Cc2ccc(F)cc2)C(=O)c2ccccc21",
    "1f": "COc1ccc(C=C2C(=O)c3ccccc3C2=O)cc1",
    "1g": "COc1ccc(C=C2C(=O)c3ccccc3C2=O)cc1OC",
    "1h": "O=C1C(=Cc2ccc(OC(F)(F)F)cc2)C(=O)c2ccccc21",
    "1i": "N#Cc1ccc(C=C2C(=O)c3ccccc3C2=O)cc1",
    "1k": "COc1cccc(C=C2C(=O)c3ccccc3C2=O)c1",
    "2a": "O=C1c2ccccc2C(=O)C12CC2c1ccccc1",
    "2b": "Cc1ccc(C2CC23C(=O)c2ccccc2C3=O)cc1",
    "2c": "Cc1ccccc1C1CC12C(=O)c1ccccc1C2=O",
    "2d": "O=C1c2ccccc2C(=O)C12CC2c1ccc(F)cc1",
    "2e": "O=C1c2ccccc2C(=O)C12CC2c1ccc(Br)cc1",
    "2f": "COc1ccc(C2CC23C(=O)c2ccccc2C3=O)cc1",
    "2g": "COc1ccc(C2CC23C(=O)c2ccccc2C3=O)cc1OC",
    "2h": "O=C1c2ccccc2C(=O)C12CC2c1ccc(OC(F)(F)F)cc1",
    "2i": "N#Cc1ccc(C2CC23C(=O)c2ccccc2C3=O)cc1",
    "2j": "O=C1c2ccccc2C(=O)C12CC2/C=C/c1ccccc1",
    "2k": "COc1cccc(C2CC23C(=O)c2ccccc2C3=O)c1",
    "2l": "O=C1c2ccccc2C(=O)C12CC2c1ccc([N+](=O)[O-])cc1",
    "2m": "CC1(C)OC(=O)C2(C[C@H]2c2ccc(C#N)cc2)C(=O)O1",
    "3e": "O=C(c1ccccc1)[C@@H]1[C@@H](c2ccc(Br)cc2)CC12C(=O)c1ccccc1C2=O",
    "3v": "CCOC(=O)C1[C@H](c2ccc(C)cc2)[C@@H](C(=O)OCC)C12C(=O)c1ccccc1C2=O",
}

YLIDE_SMILES = {
    "4a": "C[S+](C)[C-]C(=O)c1ccccc1",
    "4b": "C[S+](C)[C-]C(=O)c1ccccc1OC",
    "4c": "C[S+](C)[C-]C(=O)c1ccc(OC)cc1",
    "4d": "C[S+](C)[C-]C(=O)c1ccc(Cl)cc1",
    "4e": "C[S+](C)[C-]C(=O)c1cccs1",
    "4g": "CCOC(=O)[C-][S+](C)C",
}

METHOD_TEMPLATES = (
    (
        "2g, 2h, 2k",
        "Cyclopropane {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol), trimethylsulfoxonium iodide "
        "({Sulfoxonium.mg} mg, {Sulfoxonium.mmol} mmol) and NaH (60% dispersion in mineral oil; "
        "{NaH.mg} mg, {NaH.mmol} mmol) in DMF ({Solvent_DMF.ml} mL) according to the general "
        "procedure. Yield {Product.mg} mg ({Product.yield.percent}); {Product.appearance}; "
        "mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "3a-3i, 3k-3o",
        "Cyclobutane {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol) and {Reagent_2.name} "
        "({Reagent_2.mg} mg, {Reagent_2.mmol} mmol) in DCE ({Solvent_DCE.ml} mL) according to "
        "General procedure A. Yield {Product.mg} mg ({Product.yield.percent}); "
        "{Product.appearance}; mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "3j",
        "Cyclobutane {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol) and {Reagent_2.name} "
        "({Reagent_2.mg} mg, {Reagent_2.mmol} mmol) in DMF ({Solvent_DMF.ml} mL) according to "
        "modified General procedure A. Yield {Product.mg} mg ({Product.yield.percent}); "
        "{Product.appearance}; mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "3p, 3r-3u",
        "Cyclobutane {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol) and {Reagent_2.name} "
        "({Reagent_2.mg} mg, {Reagent_2.mmol} mmol) in DCM ({Solvent_DCM.ml} mL) according to "
        "General procedure B. Yield {Product.mg} mg "
        "({Product.yield.percent}); {Product.appearance}; mp {Product.mp} °C. "
        "Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "3q",
        "Cyclobutane {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol) and {Reagent_2.name} "
        "({Reagent_2.mg} mg, {Reagent_2.mmol} mmol) in DCE ({Solvent_DCE.ml} mL) according to "
        "General procedure B. Yield {Product.mg} mg "
        "({Product.yield.percent}); {Product.appearance}; mp {Product.mp} °C. "
        "Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "3v",
        "Cyclobutane {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol) and {Reagent_2.name} "
        "({Reagent_2.mg} mg, {Reagent_2.mmol} mmol) in DCM ({Solvent_DCM.ml} mL) according to "
        "General procedure D. Yield {Product.mg} mg ({Product.yield.percent}); "
        "{Product.appearance}; mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "3w",
        "Cyclobutane {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol) and {Reagent_2.name} "
        "({Reagent_2.mg} mg, {Reagent_2.mmol} mmol) in DCE ({Solvent_DCE.ml} mL) according to "
        "General procedure D. Yield {Product.mg} mg ({Product.yield.percent}); "
        "{Product.appearance}; mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "5",
        "Compound {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol), (4-chlorophenyl)boronic acid "
        "({Boronic_acid.mg} mg, {Boronic_acid.mmol} mmol), Na2CO3 ({Na2CO3.mg} mg, "
        "{Na2CO3.mmol} mmol) and Pd(dppf)Cl2 ({Pd_dppf_Cl2.mg} mg, {Pd_dppf_Cl2.mmol} mmol) "
        "in dioxane ({Solvent_dioxane.ml} mL) and water ({Solvent_water.ml} mL). "
        "Yield {Product.mg} mg ({Product.yield.percent}); {Product.appearance}; "
        "mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
    (
        "6",
        "Compound {Product.number} was obtained from {Reagent_1.name} "
        "({Reagent_1.mg} mg, {Reagent_1.mmol} mmol), phenylhydrazine "
        "({Phenylhydrazine.mg} mg, {Phenylhydrazine.mmol} mmol) and acetic acid "
        "({AcOH.mg} mg, {AcOH.mmol} mmol) in ethanol ({Solvent_EtOH.ml} mL). "
        "Yield {Product.mg} mg ({Product.yield.percent}); {Product.appearance}; "
        "mp {Product.mp} °C. Rf = {Product.rf.value} ({Product.rf.system}).",
    ),
)

REACTION_SCHEMAS = (
    (
        "2g, 2h, 2k",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Sulfoxonium", "1.3", "220.07", "", ""),
            ("NaH", "1.1", "40.00", "", ""),
            ("Solvent_DMF", "", "", "", "0.2"),
        ),
    ),
    (
        "3a-3i, 3k-3o",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Reagent_2", "1.5", "", "", ""),
            ("Solvent_DCE", "", "", "", "0.5"),
        ),
    ),
    (
        "3j",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Reagent_2", "2.4", "", "", ""),
            ("Solvent_DMF", "", "", "", "0.2"),
        ),
    ),
    (
        "3p, 3r-3u",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Reagent_2", "1.5", "", "", ""),
            ("Solvent_DCM", "", "", "", "0.5"),
        ),
    ),
    (
        "3q",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Reagent_2", "1.5", "", "", ""),
            ("Solvent_DCE", "", "", "", "0.5"),
        ),
    ),
    (
        "3v",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Reagent_2", "2.5", "", "", ""),
            ("Solvent_DCM", "", "", "", "0.5"),
        ),
    ),
    (
        "3w",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Reagent_2", "2.5", "", "", ""),
            ("Solvent_DCE", "", "", "", "0.5"),
        ),
    ),
    (
        "5",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Boronic_acid", "1.2", "156.38", "", ""),
            ("Na2CO3", "3", "105.99", "", ""),
            ("Pd_dppf_Cl2", "0.02", "731.70", "", ""),
            ("Solvent_dioxane", "", "", "", "1"),
            ("Solvent_water", "", "", "", "1"),
        ),
    ),
    (
        "6",
        (
            ("Reagent_1", "1", "", "", ""),
            ("Phenylhydrazine", "1.1", "108.14", "", ""),
            ("AcOH", "0.33", "60.05", "", ""),
            ("Solvent_EtOH", "", "", "", "0.3"),
        ),
    ),
)


def clean(text: str) -> str:
    return " ".join(text.split())


def set_normal(document: Document) -> None:
    for style in ("Normal", "Title", "Heading 1", "Heading 2"):
        try:
            style_object = document.styles[style]
        except KeyError:
            continue
        style_object.font.name = "Times New Roman"
        style_object.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(14)
    document.styles["Normal"].paragraph_format.space_after = Pt(0)
    document.styles["Normal"].paragraph_format.line_spacing = 1.0
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _normalize_table_ole_styles(document: Document) -> None:
    for shape in document._element.xpath(
        './/*[local-name()="tbl"]//*[local-name()="object"]//*[local-name()="shape"]'
    ):
        style = str(shape.get("style") or "")
        width = re.search(r"(?:^|;)width:([^;]+)", style)
        height = re.search(r"(?:^|;)height:([^;]+)", style)
        if width and height:
            shape.set("style", f"width:{width.group(1)};height:{height.group(1)}")


def make_si_template(path: Path) -> None:
    document = Document(REPO / "src/si_generator/templates/SI_template.docx")
    for paragraph in list(document.paragraphs):
        if "{Reagent_" in paragraph.text or paragraph.text == "{reaction.loadings}":
            paragraph._p.getparent().remove(paragraph._p)
    first_paragraph = document.paragraphs[0]._p
    definitions: list[tuple[str, bool]] = []
    for selector, template in METHOD_TEMPLATES:
        definitions.append((f"[AUTO SI: METHOD {selector}]", True))
        definitions.append((template, False))
    definitions.append(("[AUTO SI: COMPOUND TEMPLATE]", True))
    for text, bold in definitions:
        paragraph = document.add_paragraph(text)
        if paragraph.runs:
            paragraph.runs[0].bold = bold
        first_paragraph.addprevious(paragraph._p)
    set_normal(document)
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
    structure_paragraph = next(
        paragraph for paragraph in document.paragraphs if "{Product.structure}" in paragraph.text
    )
    structure_paragraph.paragraph_format.line_spacing = Pt(1)
    for run in structure_paragraph.runs:
        run.font.size = Pt(1)
    document.save(path)


def make_crystal_template(path: Path) -> None:
    document = Document()
    set_normal(document)
    paragraph = document.add_paragraph("{Product.name} ({Product.number})")
    paragraph.runs[0].bold = True
    paragraph.paragraph_format.keep_with_next = True
    for text in ("{Crystal.description}", "{Crystal.table}", "{Crystal.geometry}"):
        document.add_paragraph(text)
    document.save(path)


def _source_starts(paragraphs) -> dict[str, int]:
    starts: dict[str, int] = {}
    allowed = set(COMPOUND_NUMBERS)
    for index, paragraph in enumerate(paragraphs[:268]):
        match = re.search(r"\((\d+[a-z]?)\)\s*$", clean(paragraph.text))
        if match and match.group(1) in allowed:
            starts.setdefault(match.group(1), index)
    missing = [number for number in COMPOUND_NUMBERS if number not in starts]
    if missing:
        raise ValueError(f"Source SI is missing compound headings: {', '.join(missing)}")
    return starts


def _compound_block(paragraphs, number: str, start: int, end: int) -> dict[str, object]:
    heading = clean(paragraphs[start].text)
    name = re.sub(r"^Synthesis of\s+", "", heading, flags=re.IGNORECASE)
    name = re.sub(r"\s*\(" + re.escape(number) + r"\)\s*$", "", name)

    structure_index = SPECTRAL_STRUCTURE_PARAGRAPHS.get(number)
    if structure_index is None:
        structure_index = next(
            index
            for index in range(start + 1, end)
            if paragraphs[index]._p.xpath(".//w:object")
        )
    first_nmr = next(
        index
        for index in range(start + 1, end)
        if re.match(r"^(?:1H|13C|19F).*NMR", clean(paragraphs[index].text), flags=re.IGNORECASE)
    )
    preparation_parts = [clean(paragraphs[index].text) for index in range(start + 1, first_nmr)]
    preparation_parts = [part for part in preparation_parts if part]
    preparation = "\n\n".join(preparation_parts)

    analytical = [clean(paragraphs[index].text) for index in range(first_nmr, end)]
    analytical = [text for text in analytical if text]
    hrms = next((text for text in analytical if text.startswith("HRMS")), "")
    ir = next((text for text in analytical if text.startswith("IR")), "")
    analysis = next((text for text in analytical if text.startswith("Anal.")), "")
    found_match = re.search(r"Found\s*:?[\sСС]*([\d.]+)", hrms)
    if not found_match:
        found_match = re.search(r"Found\s*:?[\s]*([\d.]+)", hrms)
    adduct_match = re.search(r"\[M[^]]+\]\+", hrms)

    nmr_lines = [text for text in analytical if re.match(r"^(?:1H|13C|19F).*NMR", text, flags=re.IGNORECASE)]
    main_h1_used = False
    main_c13_used = False
    extras: list[str] = []
    for line in nmr_lines:
        if re.match(r"^1H", line, flags=re.IGNORECASE) and not main_h1_used:
            main_h1_used = True
            continue
        if re.match(r"^13C", line, flags=re.IGNORECASE) and not main_c13_used:
            main_c13_used = True
            continue
        extras.append(line)

    analysis_found = "-"
    if "Found:" in analysis:
        analysis_found = analysis.split("Found:", 1)[1].strip().translate(str.maketrans({"С": "C", "Н": "H"}))

    precursor = {"2g": "1g", "2h": "1h", "2k": "1k", "5": "3e", "6": "3v"}.get(number)
    if precursor is None:
        precursor_match = re.search(r"\b(?:cyclopropane|alkene)\s+(2[a-z]|1[a-z])\b", preparation)
        if not precursor_match:
            raise ValueError(f"Could not determine the limiting precursor for {number}")
        precursor = precursor_match.group(1)
    precursor_pattern = re.compile(
        rf"\b{re.escape(precursor)}\s*\(([\d.]+)\s*(mg|g),\s*[\d.]+\s*mmol",
        flags=re.IGNORECASE,
    )
    precursor_amount = precursor_pattern.search(preparation)
    if not precursor_amount:
        raise ValueError(f"Could not determine the limiting-precursor mass for {number}")
    precursor_mass_mg = float(precursor_amount.group(1)) * (
        1000 if precursor_amount.group(2).lower() == "g" else 1
    )
    yield_match = re.search(
        r"(?:Yield|yielding)\s+([\d.]+)\s*mg\s*\(([\d.]+)%\)",
        preparation,
        flags=re.IGNORECASE,
    )
    if not yield_match:
        raise ValueError(f"Could not determine product mass and yield for {number}")

    appearance_match = re.search(
        r"\b((?:(?:light|pale|dark)[- ]?)?(?:white|yellow|orange|brown|green|colorless|cream|yellowish)"
        r"(?:[- ][a-z]+)?\s+(?:solid|oil|powder|foam))\b",
        preparation,
        flags=re.IGNORECASE,
    )
    melting_point_match = re.search(
        r"\b(?:m\.?p\.?|mp)\s*=?\s*([\d]+(?:\s*[\u2013-]\s*[\d]+)?)\s*°?\s*C",
        preparation,
        flags=re.IGNORECASE,
    )
    rf_match = re.search(
        r"\bRf\s*=\s*([\d.]+(?:\s+and\s+[\d.]+)?)\.?\s*\(([^)]+)\)",
        preparation,
        flags=re.IGNORECASE,
    )
    reagent_2_match = re.search(r"\b(4[a-g])\b", preparation, flags=re.IGNORECASE)

    return {
        "number": number,
        "name": name,
        "structure_index": structure_index,
        "preparation": preparation,
        "appearance": appearance_match.group(1) if appearance_match else "",
        "melting_point": melting_point_match.group(1).replace(" ", "") if melting_point_match else "",
        "rf": f"{rf_match.group(1).rstrip('.')} ({rf_match.group(2)})" if rf_match else "",
        "hrms_found": found_match.group(1) if found_match else "",
        "hrms_adduct": adduct_match.group(0) if adduct_match else "[M+H]+",
        "ir": ir,
        "analysis": analysis_found,
        "extra_nmr": "\n".join(extras),
        "precursor": precursor,
        "precursor_smiles": PRECURSOR_SMILES[precursor],
        "reagent_2": reagent_2_match.group(1).lower() if reagent_2_match else "",
        "precursor_mass_mg": precursor_mass_mg,
        "product_mass_mg": float(yield_match.group(1)),
        "reported_yield_percent": float(yield_match.group(2)),
    }


def make_compound_table(source_path: Path, output: Path) -> list[dict[str, object]]:
    document = Document(source_path)
    paragraphs = list(document.paragraphs)
    starts = _source_starts(paragraphs)
    ordered = sorted(starts.items(), key=lambda item: item[1])
    blocks = []
    for position, (number, start) in enumerate(ordered):
        end = ordered[position + 1][1] if position + 1 < len(ordered) else 268
        blocks.append(_compound_block(paragraphs, number, start, end))

    document._body.clear_content()
    document.part._styles_part._element = deepcopy(Document().styles.element)
    set_normal(document)
    section = document.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.5)
    section.top_margin = section.bottom_margin = Inches(0.5)
    for paragraph in section.header.paragraphs + section.footer.paragraphs:
        paragraph.clear()

    headers = (
        "number",
        "structure",
        "color",
        "mp",
        "Rf",
        "HRMS",
        "Elemental_analysis",
    )
    widths = [Inches(value) for value in (0.55, 2.2, 1.1, 0.8, 2.5, 1.25, 2.2)]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    for block in blocks:
        row = table.add_row()
        values = (
            block["number"],
            "",
            block["appearance"],
            block["melting_point"],
            block["rf"],
            f"{block['hrms_adduct']} {block['hrms_found']}".strip(),
            block["analysis"],
        )
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
        structure = deepcopy(paragraphs[int(block["structure_index"])]._p.xpath(".//w:object")[0])
        row.cells[1].paragraphs[0].add_run()._r.append(structure)

    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column_index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=45, start=55, bottom=45, end=55)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = False
                if row_index == 0 or column_index in {0, 2, 3, 5}:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.font.hidden = False
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row_index == 0

    final = document.add_section(WD_SECTION.NEW_PAGE)
    final.orientation = WD_ORIENT.PORTRAIT
    final.page_width, final.page_height = Inches(8.27), Inches(11.69)
    final.left_margin = final.right_margin = Inches(0.7)
    final.top_margin = final.bottom_margin = Inches(0.7)
    _normalize_table_ole_styles(document)
    document.save(output)

    metadata = extract_structure_metadata_by_row(output)
    for row_index, block in enumerate(blocks, start=2):
        formula = metadata.get(row_index)
        if formula is None or not formula.formula:
            raise ValueError(f"Could not extract formula from structure for {block['number']}")
        block["formula"] = formula.formula
        block["smiles"] = formula.smiles
    return blocks


def make_reaction_schema(path: Path) -> None:
    document = Document()
    set_normal(document)
    headers = ("Reagents", "equiv.", "MW, g/mol", "Density, g/ml", "Concentration, M")
    for block_index, (selector, rows) in enumerate(REACTION_SCHEMAS):
        marker = document.add_paragraph(f"[AUTO SI: REACTION {selector}]")
        marker.runs[0].bold = True
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for values in rows:
            row = table.add_row()
            for cell, text in zip(row.cells, values):
                cell.text = text
        if block_index + 1 < len(REACTION_SCHEMAS):
            document.add_paragraph()
    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                _set_cell_margins(cell, top=40, start=55, bottom=40, end=55)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)
                        run.bold = row_index == 0
    document.save(path)


def _insert_reagent_oles(scope_path: Path, blocks: list[dict[str, object]]) -> None:
    import pythoncom
    import win32com.client as win32
    from rdkit import Chem
    from rdkit.Chem import AllChem

    def paste_chemdraw_object(chemical_document, cell_range) -> None:
        for attempt in range(5):
            chemical_document.Objects.Copy()
            time.sleep(0.15 + attempt * 0.15)
            try:
                cell_range.PasteSpecial(DataType=0)
                return
            except Exception:
                if attempt == 4:
                    raise

    with TemporaryDirectory(prefix="asg_scope_reagents_") as temporary:
        molecule_paths: dict[str, Path] = {}
        for block in blocks:
            precursor = str(block["precursor"])
            if precursor in molecule_paths:
                continue
            molecule = Chem.MolFromSmiles(str(block["precursor_smiles"]))
            if molecule is None:
                raise ValueError(f"Could not parse precursor structure {precursor}")
            AllChem.Compute2DCoords(molecule)
            molecule_path = Path(temporary) / f"{precursor}.mol"
            Chem.MolToMolFile(molecule, str(molecule_path))
            molecule_paths[precursor] = molecule_path
        for block in blocks:
            reagent_2 = str(block["reagent_2"])
            if not reagent_2 or reagent_2 in molecule_paths:
                continue
            molecule = Chem.MolFromSmiles(YLIDE_SMILES[reagent_2])
            if molecule is None:
                raise ValueError(f"Could not parse second-reagent structure {reagent_2}")
            AllChem.Compute2DCoords(molecule)
            molecule_path = Path(temporary) / f"{reagent_2}.mol"
            Chem.MolToMolFile(molecule, str(molecule_path))
            molecule_paths[reagent_2] = molecule_path
        for block in blocks:
            if str(block["number"]) != "3v":
                continue
            molecule = Chem.MolFromSmiles(str(block["smiles"]))
            if molecule is None:
                raise ValueError("Could not parse the single-product structure for 3v")
            AllChem.Compute2DCoords(molecule)
            molecule_path = Path(temporary) / "product_3v.mol"
            Chem.MolToMolFile(molecule, str(molecule_path))
            molecule_paths["product_3v"] = molecule_path

        pythoncom.CoInitialize()
        word = chem_draw = word_document = None
        try:
            chem_draw = win32.DispatchEx("ChemDraw.Application")
            chem_draw.Visible = False
            for precursor, molecule_path in list(molecule_paths.items()):
                chemical_document = chem_draw.Documents.Open(str(molecule_path))
                try:
                    cdx_path = molecule_path.with_suffix(".cdx")
                    chemical_document.SaveAs(str(cdx_path))
                    molecule_paths[precursor] = cdx_path
                finally:
                    chemical_document.Close(False)
            word = win32.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            word_document = word.Documents.Open(str(scope_path.resolve()))
            table = word_document.Tables(1)
            for row_index, block in enumerate(blocks, start=2):
                chem_document = chem_draw.Documents.Open(str(molecule_paths[str(block["precursor"])]))
                try:
                    cell_range = table.Cell(row_index, 1).Range
                    cell_range.End -= 1
                    cell_range.Text = ""
                    paste_chemdraw_object(chem_document, cell_range)
                finally:
                    chem_document.Close(False)
                reagent_2 = str(block["reagent_2"])
                if reagent_2:
                    chem_document = chem_draw.Documents.Open(str(molecule_paths[reagent_2]))
                    try:
                        cell_range = table.Cell(row_index, 3).Range
                        cell_range.End -= 1
                        cell_range.Text = ""
                        paste_chemdraw_object(chem_document, cell_range)
                    finally:
                        chem_document.Close(False)
                if row_index % 6 == 0:
                    word_document.Save()
            for row_index, block in enumerate(blocks, start=2):
                if str(block["number"]) != "3v":
                    continue
                chem_document = chem_draw.Documents.Open(str(molecule_paths["product_3v"]))
                try:
                    cell_range = table.Cell(row_index, 5).Range
                    cell_range.End -= 1
                    cell_range.Text = ""
                    paste_chemdraw_object(chem_document, cell_range)
                finally:
                    chem_document.Close(False)
            word_document.Save()
        finally:
            if word_document is not None:
                word_document.Close(False)
            if word is not None:
                word.Quit()
            if chem_draw is not None:
                try:
                    if chem_draw.Documents.Count == 0:
                        chem_draw.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()


def make_scope(source_path: Path, blocks: list[dict[str, object]], output: Path) -> None:
    document = Document(source_path)
    source_paragraphs = list(document.paragraphs)
    document._body.clear_content()
    set_normal(document)
    section = document.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.5)
    section.top_margin = section.bottom_margin = Inches(0.5)
    headers = (
        "Reagent_1",
        "Mass of Reagent_1, mg",
        "Reagent_2",
        "Mass of Reagent_2, mg",
        "Product",
        "Product_number",
        "Mass of product, mg",
    )
    widths = [Inches(value) for value in (1.65, 1.15, 1.65, 1.15, 1.9, 0.9, 1.2)]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for block in blocks:
        row = table.add_row()
        values = (
            "",
            f"{float(block['precursor_mass_mg']):g}",
            "",
            "",
            "",
            str(block["number"]),
            f"{float(block['product_mass_mg']):g}",
        )
        for cell, value in zip(row.cells, values):
            cell.text = value
        if str(block["number"]) != "3v":
            product = deepcopy(source_paragraphs[int(block["structure_index"])]._p.xpath(".//w:object")[0])
            row.cells[4].paragraphs[0].add_run()._r.append(product)

    for row_index, row in enumerate(table.rows):
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
    document.save(output)
    _insert_reagent_oles(output, blocks)
    normalized = Document(output)
    _normalize_table_ole_styles(normalized)
    normalized.save(output)

    metadata = extract_structure_metadata_by_cell(output)
    for row_index, block in enumerate(blocks, start=2):
        reagent = metadata.get((1, row_index, 1))
        reagent_2 = metadata.get((1, row_index, 3))
        product = metadata.get((1, row_index, 5))
        if reagent is None or not reagent.formula:
            raise ValueError(f"Scope is missing the precursor structure for {block['number']}")
        if block["reagent_2"] and (reagent_2 is None or not reagent_2.formula):
            raise ValueError(f"Scope is missing Reagent_2 for {block['number']}")
        if product is None or product.formula != block.get("formula"):
            raise ValueError(f"Scope product structure differs from the compound table for {block['number']}")


def _write_spectra_archives(source_path: Path, main_path: Path, additional_path: Path) -> dict[str, object]:
    coverage: dict[str, dict[str, list[str]]] = {
        number: {"primary": [], "additional": []} for number in COMPOUND_NUMBERS
    }
    with (
        zipfile.ZipFile(source_path) as source,
        zipfile.ZipFile(main_path, "w", zipfile.ZIP_DEFLATED) as main,
        zipfile.ZipFile(additional_path, "w", zipfile.ZIP_DEFLATED) as additional,
    ):
        for item in source.infolist():
            parts = item.filename.replace("\\", "/").split("/")
            if item.is_dir() or len(parts) < 4 or parts[0] != "fid" or parts[1] not in coverage:
                continue
            number, experiment = parts[1], parts[2]
            is_additional = (
                experiment == "19F"
                or "minor diastereomer" in experiment
                or (number == "3a" and experiment == "13C")
            )
            normalized_experiment = experiment.replace(" (mixture of diastereomers)", "")
            target_parts = [number, normalized_experiment, *parts[3:]]
            target = "/".join(target_parts)
            archive = additional if is_additional else main
            archive.writestr(target, source.read(item))
            bucket = "additional" if is_additional else "primary"
            if normalized_experiment not in coverage[number][bucket]:
                coverage[number][bucket].append(normalized_experiment)

    for number, spectra in coverage.items():
        primary = set(spectra["primary"])
        if not any(value.startswith("1H") for value in primary):
            raise ValueError(f"{number}: primary 1H spectrum is missing")
        if not any(value.startswith("13C") for value in primary):
            raise ValueError(f"{number}: primary 13C spectrum is missing")
    return coverage


def _write_crystallography(source: Document, cif_root: Path, root: Path) -> dict[str, object]:
    result: dict[str, object] = {}
    for number, (image_index, growth_index, ccdc, filename, pdf_name) in CRYSTAL_RECORDS.items():
        source_cif = cif_root / filename
        folder = root / "CIF_source" / number
        folder.mkdir(parents=True, exist_ok=True)
        target_cif = folder / f"{number}.cif"
        shutil.copy2(source_cif, target_cif)
        relationship_ids = source.paragraphs[image_index]._p.xpath(".//a:blip/@r:embed")
        if len(relationship_ids) != 1:
            raise ValueError(f"Expected one original ORTEP image for {number}")
        (folder / "ORTEP.png").write_bytes(source.part.related_parts[relationship_ids[0]].blob)
        caption = "ORTEP drawing with thermal ellipsoids at the 50% probability level."
        if number == "3v":
            caption = "Minor diastereomer B. " + caption
        metadata = {
            "ccdc": ccdc,
            "growth": clean(source.paragraphs[growth_index].text),
            "image": "ORTEP.png",
            "caption": caption,
            "include_geometry": False,
        }
        (folder / f"{number}.json").write_text(
            json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        if pdf_name:
            shutil.copy2(source_cif.parent / pdf_name, folder / f"{number}.checkcif.pdf")
        result[number] = {
            "ccdc": ccdc,
            "cif_source": filename,
            "cif_sha256": hashlib.sha256(source_cif.read_bytes()).hexdigest(),
            "ortep_paragraph": image_index,
            "growth_paragraph": growth_index,
        }
    return result


def _strip_word_comments(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        entries = {item.filename: source.read(item.filename) for item in source.infolist()}
    for name in list(entries):
        if name.startswith("word/comments"):
            entries.pop(name)
            continue
        if not name.endswith((".xml", ".rels")):
            continue
        try:
            root = etree.fromstring(entries[name])
        except etree.XMLSyntaxError:
            continue
        changed = False
        for element in list(root.iter()):
            local_name = etree.QName(element).localname
            should_remove = local_name in {"commentRangeStart", "commentRangeEnd", "commentReference"}
            if local_name == "Relationship" and "comments" in str(element.get("Target") or ""):
                should_remove = True
            if local_name == "Override" and "comments" in str(element.get("PartName") or ""):
                should_remove = True
            if should_remove and element.getparent() is not None:
                element.getparent().remove(element)
                changed = True
        if changed:
            entries[name] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
    temporary = path.with_suffix(".comments-cleaned.docx")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as output:
        for name, payload in entries.items():
            output.writestr(name, payload)
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-docx", required=True, type=Path)
    parser.add_argument("--fid-zip", required=True, type=Path)
    parser.add_argument("--cif-root", required=True, type=Path)
    parser.add_argument("--source-cif", action="append", default=[], type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    source = Document(args.source_docx)
    make_crystal_template(root / "Crystallography_template.docx")
    blocks = make_compound_table(args.source_docx, root / "Compound_table.docx")
    make_si_template(root / "SI_template.docx")
    make_reaction_schema(root / "Reaction_schema.docx")
    make_scope(args.source_docx, blocks, root / "Scope.docx")
    build_unified_input_docx(
        root / "Compound_table.docx",
        root / "All_in_one_input.docx",
        reaction_schema=root / "Reaction_schema.docx",
        scope=root / "Scope.docx",
        si_template=root / "SI_template.docx",
        crystallography_template=root / "Crystallography_template.docx",
    )
    for filename in (
        "Compound_table.docx",
        "Reaction_schema.docx",
        "Scope.docx",
        "SI_template.docx",
        "Crystallography_template.docx",
        "All_in_one_input.docx",
    ):
        _strip_word_comments(root / filename)

    coverage = _write_spectra_archives(
        args.fid_zip,
        root / "Spectra_source.zip",
        root / "Additional_spectra.zip",
    )
    crystals = _write_crystallography(source, args.cif_root, root)
    source_cif_root = root / "Source_files" / "CIF"
    if args.source_cif:
        source_cif_root.mkdir(parents=True, exist_ok=True)
        for source_cif in args.source_cif:
            shutil.copy2(source_cif, source_cif_root / source_cif.name)
    additional_source_files = [
        {
            "path": str(path.relative_to(root)).replace("\\", "/"),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        }
        for path in sorted(source_cif_root.glob("*.cif"))
    ] if source_cif_root.exists() else []
    provenance = {
        "source_document": args.source_docx.name,
        "source_sha256": hashlib.sha256(args.source_docx.read_bytes()).hexdigest(),
        "spectra_archive": args.fid_zip.name,
        "spectra_sha256": hashlib.sha256(args.fid_zip.read_bytes()).hexdigest(),
        "additional_source_files": additional_source_files,
        "compound_order": [str(block["number"]) for block in blocks],
        "spectra": coverage,
        "crystallography": crystals,
    }
    (root / "provenance.json").write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(root)


if __name__ == "__main__":
    main()
